# -*- coding: utf-8 -*-
"""Build the first 10-minute presentation, keeping every deliverable under v2.

Python dependencies: python-pptx, matplotlib, numpy, Pillow.
Equations are high-resolution transparent images; all body text and most diagrams
remain native, editable PowerPoint objects. Formula source is kept in this script.
"""
import re
from pathlib import Path
import json
import math
import os
import sys
import tempfile
LOCAL_DEPS = Path(tempfile.gettempdir()) / 'codex-dueling-ppt-deps'
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import mathtext
from matplotlib.font_manager import FontProperties
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
PREVIEW = ROOT / 'preview'
ASSETS.mkdir(exist_ok=True)
PREVIEW.mkdir(exist_ok=True)
OUTFILE = ROOT / 'Dueling_Bandits_Sections_1-4_v2.pptx'
(W, H) = (13.333333, 7.5)
TOTAL_SLIDES = 13
FONT = 'Calibri'
C = {'bg': 'F7F8FC', 'ink': '16263B', 'muted': '637187', 'line': 'DCE3EB', 'white': 'FFFFFF', 'teal': '007E87', 'teal_light': 'E5F2F1', 'amber': 'B57725', 'amber_light': 'FFF1DD', 'blue': '4B62B5', 'blue_light': 'EAEFFA', 'cyan': '087EA7', 'cyan_light': 'E6F3F8', 'navy': '13273E', 'pale': 'B4C6D5', 'green': '3D875C'}
SECTION = {1: ('PROBLEM SETUP', 'Problem Setup', 'teal', 'teal_light'), 2: ('ASSUMPTIONS', 'Assumptions', 'amber', 'amber_light'), 3: ('ALGORITHM', 'Algorithm', 'blue', 'blue_light'), 4: ('CONFIDENCE BOUND', 'Confidence', 'cyan', 'cyan_light')}
matplotlib.rcParams.update({'mathtext.fontset': 'stix', 'font.family': 'DejaVu Sans', 'axes.spines.top': False, 'axes.spines.right': False, 'axes.labelcolor': '#637187', 'xtick.color': '#637187', 'ytick.color': '#637187', 'axes.edgecolor': '#BCC8D2', 'figure.facecolor': 'none', 'savefig.facecolor': 'none'})
prs = Presentation()
(prs.slide_width, prs.slide_height) = (Inches(W), Inches(H))
prs.core_properties.title = 'Contextual Dueling Bandits: Setup, Assumptions, Algorithm and Confidence'
prs.core_properties.subject = 'A Clean Proof for Dueling Bandits, Sections 1–4, a 10-minute presentation'
prs.core_properties.author = ''
prs.core_properties.keywords = 'Dueling Bandits, Ridge-MLE, UCB, Confidence Bound'
prs.core_properties.comments = 'v2: English edition, four balanced sections, with speaker notes.'
DECK = []
SPEAKER_NOTES = json.loads((ROOT / 'speaker_notes_source.json').read_text(encoding='utf-8'))
FORMULAS = []

def rgb(value):
    value = C.get(value, value)
    return RGBColor.from_string(value)

def rect(slide, x, y, w, h, fill='white', line=None, radius=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if radius:
        shp.adjustments[0] = 0.1
    return shp

def text(slide, value, x, y, w, h, size=22, color='ink', bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, font=FONT):
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for (i, line) in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(4)
        p.line_spacing = 1.12
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        rpr = run._r.get_or_add_rPr()
        rpr.set('lang', 'en-US')
        ea = OxmlElement('a:ea')
        ea.set('typeface', font)
        rpr.append(ea)
    return shp

def line(slide, x1, y1, x2, y2, color='line', width=1.4, arrow=False):
    shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width)
    if arrow:
        end = OxmlElement('a:tailEnd')
        end.set('type', 'triangle')
        end.set('w', 'med')
        end.set('len', 'med')
        shp.line._get_or_add_ln().append(end)
    return shp

def dot(slide, x, y, r=0.06, color='teal'):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    return shp

def equation(slide, latex, x, y, maxw, size=26, color='ink', maxh=None, center=False):
    """Render at the requested point size, shrinking only when necessary."""
    latex = latex.replace('\\|', '\\Vert ')
    idx = len(FORMULAS) + 1
    path = ASSETS / f'equation-{idx:03d}.png'
    with matplotlib.rc_context({'text.color': '#' + C.get(color, color)}):
        mathtext.math_to_image('$' + latex + '$', str(path), prop=FontProperties(size=size), dpi=240)
    im = Image.open(path).convert('RGBA')
    padded = Image.new('RGBA', (im.width + 18, im.height + 18))
    padded.alpha_composite(im, (9, 9))
    padded.save(path)
    (w, h) = (padded.width / 240, padded.height / 240)
    scale = min(1.0, maxw / w, maxh / h if maxh else 1.0)
    (w, h) = (w * scale, h * scale)
    actual_x = x + (maxw - w) / 2 if center else x
    pic = slide.shapes.add_picture(str(path), Inches(actual_x), Inches(y), width=Inches(w), height=Inches(h))
    pic.name = f'Equation {idx}: {latex}'
    FORMULAS.append({'file': path.name, 'latex': latex, 'slide': len(prs.slides), 'effective_font_pt': round(size * scale, 1)})
    return pic

def pill(slide, label, x, y, w, accent='teal', light='teal_light', size=13):
    rect(slide, x, y, w, 0.34, light, radius=True)
    text(slide, label, x + 0.09, y + 0.045, w - 0.18, 0.25, size=size, color=accent, bold=True)

def footer(slide, number, section, source, seconds, dark=False):
    ink = 'pale' if dark else 'muted'
    line(slide, 0.64, 7.02, 12.69, 7.02, '426077' if dark else 'line', 0.8)
    text(slide, source, 0.67, 7.13, 9.9, 0.2, size=10, color=ink)
    text(slide, f'{seconds} sec   ·   {number:02d} / {TOTAL_SLIDES}', 11.0, 7.1, 1.65, 0.25, size=11, color=ink, align=PP_ALIGN.RIGHT)
    for i in range(1, 5):
        col = SECTION[i][2] if i == section else '426077' if dark else 'DCE3EB'
        rect(slide, 0.67 + (i - 1) * 0.53, 6.88, 0.4, 0.045, col)

def new_slide(section, title, subtitle, seconds, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb('bg')
    (english, _, accent, light) = SECTION[section]
    text(slide, f'{section:02d}  {english}', 0.68, 0.38, 11.9, 0.3, size=12, color=accent, bold=True)
    text(slide, title, 0.66, 0.94, 12.05, 0.66, size=31, bold=True)
    text(slide, subtitle, 0.69, 1.67, 11.95, 0.49, size=17, color='muted')
    footer(slide, len(prs.slides), section, source, seconds)
    return slide

def record(slide, title, section, seconds, source, key):
    speech = SPEAKER_NOTES[len(prs.slides) - 1]
    notes = speech['script']
    optional = speech.get('optional', '')
    detail = f'\n\nOptional detail (outside the timed script):\n{optional}' if optional else ''
    slide.notes_slide.notes_text_frame.text = f'Slide {len(prs.slides)} | {title}\nSuggested time: {seconds} sec\nKey point: {key}\n\n{notes}{detail}\n\n{source}'
    DECK.append({'page': len(prs.slides), 'title': title, 'section': section, 'seconds': seconds, 'key': key, 'notes': notes, 'optional': optional, 'source': source, 'spoken_words': len(re.findall("\\b[\\w'-]+\\b", notes))})

def plot_sigmoid():
    (fig, ax) = plt.subplots(figsize=(4.8, 2.15), dpi=200)
    z = np.linspace(-5, 5, 350)
    ax.plot(z, 1 / (1 + np.exp(-z)), color='#' + C['teal'], lw=3)
    ax.axhline(0.5, color='#C4CED6', ls='--', lw=1)
    ax.axvline(0, color='#C4CED6', ls='--', lw=1)
    ax.scatter([0], [0.5], color='#' + C['teal'], s=36, zorder=4)
    ax.set(xlim=(-5, 5), ylim=(-0.02, 1.03), xticks=[-4, 0, 4], yticks=[0, 0.5, 1])
    ax.tick_params(labelsize=12)
    ax.set_xlabel('utility gap $z$', fontsize=12)
    ax.set_ylabel('$\\mu(z)$', fontsize=14, rotation=0, labelpad=17)
    fig.tight_layout(pad=0.7)
    target = ASSETS / 'logistic-link.png'
    fig.savefig(target, transparent=True)
    plt.close(fig)
    return target

def plot_curvature():
    (fig, ax) = plt.subplots(figsize=(5.0, 2.6), dpi=220)
    z = np.linspace(-6, 6, 400)
    p = 1 / (1 + np.exp(-z))
    k = np.exp(-2) / (1 + np.exp(-2)) ** 2
    ax.axvspan(-2, 2, color='#F7E7CC', alpha=0.75)
    ax.plot(z, p * (1 - p), color='#' + C['amber'], lw=3)
    ax.hlines(k, -2, 2, color='#' + C['ink'], ls='--', lw=1.4)
    ax.vlines([-2, 2], 0, k, color='#BBA27F', ls=':', lw=1.2)
    ax.annotate('$\\kappa_\\mu$', (2.1, k + 0.012), fontsize=18, color='#' + C['ink'])
    ax.set(xlim=(-6, 6), ylim=(0, 0.28), xticks=[-6, -2, 0, 2, 6], yticks=[0, 0.1, 0.2])
    ax.tick_params(labelsize=12)
    ax.set_xlabel('$z$  (illustration: $S=1$)', fontsize=12)
    ax.set_ylabel("$\\mu'(z)$", fontsize=16, rotation=0, labelpad=15)
    fig.tight_layout(pad=0.7)
    target = ASSETS / 'local-curvature.png'
    fig.savefig(target, transparent=True)
    plt.close(fig)
    return target
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb('navy')
rect(s, 0.68, 0.7, 0.1, 3.1, 'teal')
text(s, 'A CLEAN PROOF FOR DUELING BANDITS', 0.98, 0.72, 11.3, 0.35, size=15, color='pale', bold=True)
text(s, 'From Pairwise Preferences\nto Confidence Bounds', 0.95, 1.42, 9.4, 1.9, size=44, color='white', bold=True)
text(s, 'Contextual Dueling Bandits', 1.0, 3.48, 10.7, 0.48, size=27, color='pale', font='Calibri')
text(s, 'Sections 1–4  ·  About 10 minutes  ·  v2', 1.0, 4.12, 10.5, 0.36, size=17, color='pale')
for (i, (cn, en, timing)) in enumerate([('Problem Setup', 'What do we observe?', '2 min 10 sec'), ('Assumptions', 'What makes it learnable?', '2 min 10 sec'), ('Algorithm', 'How do we choose?', '2 min 30 sec'), ('Confidence', 'Why trust the estimate?', '2 min 30 sec')]):
    x = 0.98 + i * 3.08
    text(s, f'0{i + 1}', x, 5.12, 0.55, 0.35, size=15, color='66C2C1', bold=True)
    text(s, cn, x, 5.6, 2.7, 0.48, size=24, color='white', bold=True)
    text(s, timing, x, 6.2, 2.7, 0.3, size=13, color='pale')
footer(s, 1, 0, 'Zhuohua LI · A Clean Proof for Dueling Bandits · February 2026 · §§1–4', 10, True)
record(s, 'From Pairwise Preferences to Confidence Bounds', 'Opening', 10, 'Source: Sections 1–4, pp.1–7', 'Preference feedback → parameter estimation → confident comparisons')
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb('bg')
text(s, 'CONTENTS  /  TALK OUTLINE', 0.68, 0.38, 11.9, 0.3, 12, 'teal', True)
text(s, 'Contents', 0.66, 0.94, 12.05, 0.66, 34, bold=True)
text(s, 'Four questions connect preference feedback to confidence-based decisions.', 0.69, 1.67, 11.95, 0.49, 18, 'muted')
agenda = [('Problem Setup: Contextual Dueling Bandits', 'What we observe: arm features, pairwise preferences and regret', 'Slides 03–04', '2 min 10 sec'), ('Assumptions', 'What we assume: boundedness and a local curvature lower bound', 'Slides 05–06', '2 min 10 sec'), ('Algorithm: Ridge-MLE Dueling UCB', 'How we choose: parameter estimation, a greedy arm and a UCB challenger', 'Slides 07–09', '2 min 30 sec'), ('Confidence Bound', 'Why it works: a parameter ellipsoid and a pairwise utility bound', 'Slides 10–12', '2 min 30 sec')]
for (i, (heading, description, pages, timing)) in enumerate(agenda, start=1):
    y = 2.35 + (i - 1) * 1.06
    accent = SECTION[i][2]
    rect(s, 0.7, y, 11.94, 0.92, 'white', radius=True)
    rect(s, 0.7, y + 0.15, 0.045, 0.62, accent)
    text(s, f'0{i}', 0.96, y + 0.23, 0.72, 0.49, 27, accent, True)
    text(s, heading, 1.89, y + 0.12, 8.48, 0.37, 21, bold=True, font='Calibri')
    text(s, description, 1.91, y + 0.56, 8.4, 0.29, 15, 'muted')
    text(s, pages, 10.62, y + 0.19, 1.68, 0.32, 17, accent, True, PP_ALIGN.RIGHT)
    text(s, timing, 10.62, y + 0.59, 1.68, 0.25, 12, 'muted', align=PP_ALIGN.RIGHT)
footer(s, len(prs.slides), 0, 'A Clean Proof for Dueling Bandits · Sections 1–4 · Outline and slide numbers', 20)
record(s, 'Contents', 'Contents', 20, 'Source: Sections 1–4, pp.1–7', 'Problem setup → assumptions → algorithm → confidence bound')
s = new_slide(1, 'Compare two arms; observe one preference', 'Contextual: arms have features, and the decision set may change each round.', 65, 'Source: Section 1, p.1; Eqs. (1)–(3)')
xs = [0.7, 4.97, 9.24]
widths = [3.4, 3.4, 3.4]
for (x, w) in zip(xs, widths):
    rect(s, x, 2.43, w, 2.73, 'white', radius=True)
for (i, label) in enumerate(['Observe the arm set', 'Select a pair', 'Observe a preference']):
    pill(s, f'{i + 1}  {label}', xs[i] + 0.22, 2.66, 2.65)
equation(s, '\\mathcal{X}_t\\subseteq\\mathbb{R}^{d}', 0.96, 3.3, 2.95, 31)
text(s, 'Each arm has features x.', 0.96, 4.17, 2.98, 0.75, 18)
equation(s, 'x_{t,1},\\ x_{t,2}\\in\\mathcal{X}_t', 5.21, 3.3, 2.98, 29)
text(s, 'Example: two recommendations.', 5.23, 4.17, 2.98, 0.75, 18)
equation(s, 'y_t\\in\\{0,1\\}', 9.5, 3.3, 2.95, 31)
text(s, 'y = 1: the first arm wins.', 9.48, 4.17, 3.0, 0.75, 18)
line(s, 4.24, 3.72, 4.8, 3.72, 'teal', 2.4, True)
line(s, 8.5, 3.72, 9.07, 3.72, 'teal', 2.4, True)
rect(s, 0.7, 5.47, 11.94, 1.06, 'teal_light', radius=True)
text(s, 'Latent utility', 0.94, 5.68, 2.75, 0.44, 18, 'teal', True)
equation(s, 'f(x)=\\langle\\theta^{\\star},x\\rangle', 3.62, 5.67, 3.27, 27)
text(s, 'Difference feature', 7.02, 5.7, 2.43, 0.4, 16, 'teal')
equation(s, '\\widetilde{x}_t=x_{t,1}-x_{t,2}', 9.1, 5.71, 3.22, 25)
record(s, 'Compare two arms; observe one preference', '1 Problem Setup', 65, 'Source: Section 1, p.1', 'Separate observed preferences from latent utilities.')
s = new_slide(1, 'Preference probabilities and utility regret', 'Win probabilities describe the feedback; regret measures the quality of our choices.', 65, 'Source: Section 1, p.1; BTL model and regret definition')
rect(s, 0.7, 2.4, 5.77, 4.11, 'white', radius=True)
rect(s, 6.68, 2.4, 5.96, 4.11, 'white', radius=True)
pill(s, 'LOGISTIC / BTL FEEDBACK', 0.94, 2.62, 2.73)
equation(s, '\\Pr(y_t=1\\mid\\mathcal{F}_{t-1})=\\mu(\\langle\\theta^{\\star},\\widetilde{x}_t\\rangle)', 0.94, 3.16, 5.22, 25)
equation(s, '\\mu(z)=\\frac{1}{1+e^{-z}}', 0.96, 3.73, 4.95, 25)
sigmoid = plot_sigmoid()
s.shapes.add_picture(str(sigmoid), Inches(1.02), Inches(4.3), width=Inches(4.96), height=Inches(1.78))
text(s, 'Equal utility: 50% wins. Higher utility: a higher win probability.', 0.98, 6.13, 5.19, 0.25, 12, 'muted')
pill(s, 'OBJECTIVE: BOTH ARMS SHOULD BE GOOD', 6.94, 2.62, 4.66)
equation(s, 'x_t^{\\star}\\in\\arg\\max_{x\\in\\mathcal{X}_t}f(x)', 6.96, 3.2, 5.1, 27)
text(s, 'Regret = sum of the two utility losses', 6.96, 4.01, 5.2, 0.37, 18)
equation(s, 'r_t=[f(x_t^{\\star})-f(x_{t,1})]', 6.97, 4.51, 5.1, 27)
equation(s, '\\quad+[f(x_t^{\\star})-f(x_{t,2})]', 6.97, 4.99, 5.1, 27)
equation(s, 'R_T=\\sum_{t=1}^{T}r_t', 6.97, 5.55, 4.9, 25)
record(s, 'Preference probabilities and utility regret', '1 Problem Setup', 65, 'Source: Section 1, p.1', 'Preferences are random feedback; regret measures latent utility loss.')
s = new_slide(2, 'Bounded features and parameters control the scale', 'These bounds control curvature, design-matrix growth and ridge bias.', 65, 'Source: Section 1, Eqs. (1)(2); Section 2, Assumption 1; pp.1–2')
rect(s, 0.7, 2.43, 5.15, 3.33, 'white', radius=True)
pill(s, 'SCALE BOUNDS', 0.96, 2.68, 2.31, 'amber', 'amber_light')
equation(s, '\\|x\\|_2\\leq 1', 0.96, 3.36, 4.3, 34)
text(s, 'Bounded features', 3.55, 3.49, 1.95, 0.4, 18, 'muted')
equation(s, '\\|\\theta^{\\star}\\|_2\\leq S', 0.96, 4.4, 4.3, 34)
text(s, 'Known radius S', 3.55, 4.53, 1.95, 0.4, 18, 'muted')
line(s, 6.02, 4.07, 6.56, 4.07, 'amber', 2.6, True)
rect(s, 6.77, 2.43, 5.87, 3.33, 'amber_light', radius=True)
text(s, 'Difference features and true scores', 7.03, 2.71, 5.17, 0.45, 21, 'amber', True)
equation(s, '\\|\\widetilde{x}_t\\|_2\\leq\\|x_{t,1}\\|_2+\\|x_{t,2}\\|_2\\leq2', 7.03, 3.48, 5.22, 27)
equation(s, '|\\langle\\theta^{\\star},\\widetilde{x}_t\\rangle|\\leq2S', 7.03, 4.45, 5.18, 31)
text(s, 'Role', 0.94, 6.07, 1.55, 0.35, 18, 'amber', True)
text(s, 'Bounded utility gaps provide a scale for the estimation error.', 2.52, 6.07, 9.76, 0.4, 21)
record(s, 'Bounded features and parameters control the scale', '2 Assumptions', 65, 'Source: Sections 1–2, pp.1–2', 'Boundedness supports control of features, curvature and ridge bias.')
s = new_slide(2, 'Local curvature makes parameter error controllable', 'The logistic curve flattens in the tails; the proof needs curvature on the relevant scores.', 65, 'Source: Section 2, Assumption 2; Section 4, Lemma 1, Step 2; pp.2, 5')
rect(s, 0.7, 2.43, 5.66, 2.93, 'white', radius=True)
curve = plot_curvature()
s.shapes.add_picture(str(curve), Inches(0.88), Inches(2.6), width=Inches(5.25), height=Inches(2.52))
rect(s, 6.57, 2.43, 6.07, 2.93, 'white', radius=True)
pill(s, 'ASSUMPTION 2', 6.83, 2.66, 2.35, 'amber', 'amber_light')
equation(s, "\\mu'(z)=\\mu(z)(1-\\mu(z))\\geq\\kappa_\\mu>0", 6.83, 3.26, 5.43, 28)
text(s, 'Intermediate scores used in the proof:', 6.84, 4.03, 5.18, 0.42, 20)
equation(s, '\\xi_{s,t}\\ \\mathrm{between}\\ \\langle\\theta^{\\star},\\widetilde{x}_s\\rangle\\ \\mathrm{and}\\ \\langle\\widehat\\theta_t,\\widetilde{x}_s\\rangle', 6.84, 4.63, 5.32, 24)
rect(s, 0.7, 5.59, 11.94, 1.04, 'amber_light', radius=True)
text(s, 'Scope', 0.94, 5.8, 1.35, 0.35, 17, 'amber', True)
text(s, 'A bound on the true parameter alone does not control the fitted scores.\nThe confidence result also needs curvature at every intermediate score.', 2.39, 5.76, 9.86, 0.72, 17)
record(s, 'Local curvature makes parameter error controllable', '2 Assumptions', 65, 'Source: Section 2, p.2; Section 4, p.5. The scope note makes a required condition explicit.', 'The curvature bound must cover scores between the true and fitted values.')
s = new_slide(3, 'Estimate utility with ridge logistic regression', 'Use the first t−1 rounds to estimate the parameter for round t.', 50, 'Source: Section 3.1, pp.2–4; Eqs. (4)–(8)')
rect(s, 0.7, 2.44, 11.94, 1.38, 'white', radius=True)
pill(s, '1  Negative log-likelihood of one duel', 0.93, 2.65, 3.55, 'blue', 'blue_light')
equation(s, '\\ell(y,z)=-y\\log\\mu(z)-(1-y)\\log(1-\\mu(z))', 4.74, 2.8, 7.49, 27)
text(s, 'Score', 0.94, 3.18, 1.21, 0.28, 14, 'muted')
equation(s, 'z=\\langle\\theta,\\widetilde{x}_s\\rangle', 2.12, 3.16, 2.2, 23)
rect(s, 0.7, 4.01, 11.94, 1.59, 'blue_light', radius=True)
pill(s, '2  Sum the losses and add ridge', 0.93, 4.24, 3.72, 'blue', 'white')
equation(s, '\\widehat\\theta_t=\\arg\\min_{\\theta\\in\\mathbb{R}^{d}}\\left\\{\\sum_{s<t}\\ell(y_s,\\langle\\theta,\\widetilde{x}_s\\rangle)+\\frac{\\lambda}{2}\\|\\theta\\|_2^2\\right\\}', 4.82, 4.38, 7.44, 30)
text(s, 'λ > 0', 0.95, 6.01, 1.31, 0.43, 22, 'blue', True)
text(s, 'Strict convexity and stable estimation, with a ridge-bias term.', 2.28, 6.02, 10.03, 0.52, 20)
record(s, 'Estimate utility with ridge logistic regression', '3 Algorithm', 50, 'Source: Section 3.1, pp.2–4', 'Bernoulli preferences yield logistic loss; ridge regularization stabilizes estimation.')
s = new_slide(3, 'The first-order condition starts the proof', 'The logistic loss derivative simplifies to predicted probability minus observed label.', 50, 'Source: Section 3.2, p.4; Eqs. (9)(10)')
for (y, tag) in [(2.44, 'Score derivative'), (3.66, 'Parameter gradient'), (5.04, 'At the minimizer')]:
    rect(s, 0.7, y, 11.94, 1.07 if y < 5 else 1.37, 'white' if y < 5 else 'blue_light', radius=True)
    text(s, tag, 0.95, y + 0.31, 2.68, 0.43, 20, 'blue', True)
equation(s, '\\frac{\\partial\\ell(y,z)}{\\partial z}=\\mu(z)-y', 4.15, 2.65, 7.86, 30)
equation(s, '\\nabla_\\theta\\ell(y_s,\\langle\\theta,\\widetilde{x}_s\\rangle)=[\\mu(\\langle\\theta,\\widetilde{x}_s\\rangle)-y_s]\\widetilde{x}_s', 4.15, 3.94, 7.93, 27)
equation(s, '\\sum_{s<t}[\\mu(\\langle\\widehat\\theta_t,\\widetilde{x}_s\\rangle)-y_s]\\widetilde{x}_s+\\lambda\\widehat\\theta_t=0', 4.15, 5.35, 7.92, 30)
record(s, 'The first-order condition starts the proof', '3 Algorithm', 50, 'Source: Section 3.2, p.4', 'Use the stationarity equation as the starting point of the error analysis.')
s = new_slide(3, 'Choose a greedy arm, then an optimistic challenger', 'Ridge-MLE Dueling UCB: exploit the current estimate, then explore through uncertainty.', 50, 'Source: Section 3, Algorithm 1, p.2; b = beta_T / kappa_mu')
rect(s, 0.7, 2.42, 8.26, 1.37, 'white', radius=True)
pill(s, 'First arm: highest estimated utility', 0.93, 2.62, 3.23, 'blue', 'blue_light')
equation(s, 'x_{t,1}\\in\\arg\\max_{x\\in\\mathcal{X}_t}\\langle\\widehat\\theta_t,x\\rangle', 4.42, 2.81, 4.15, 28)
rect(s, 0.7, 3.99, 8.26, 2.49, 'blue_light', radius=True)
pill(s, 'Second arm: largest UCB relative to the first', 0.93, 4.21, 4.54, 'blue', 'white')
equation(s, 'x_{t,2}\\in\\arg\\max_{x\\in\\mathcal{X}_t}\\{\\langle\\widehat\\theta_t,x-x_{t,1}\\rangle', 0.95, 4.94, 7.61, 29)
equation(s, '\\qquad\\qquad\\quad+b\\,\\|x-x_{t,1}\\|_{V_{t-1}^{-1}}\\}', 0.95, 5.67, 7.61, 29)
rect(s, 9.18, 2.42, 3.46, 4.06, 'white', radius=True)
text(s, 'Design matrix', 9.43, 2.67, 2.91, 0.43, 21, 'blue', True)
equation(s, 'V_0=\\frac{\\lambda}{\\kappa_\\mu}I_d', 9.43, 3.38, 2.82, 28)
text(s, 'Update after the duel', 9.43, 4.3, 2.86, 0.38, 18)
equation(s, 'V_t=V_{t-1}+\\widetilde{x}_t\\widetilde{x}_t^{\\top}', 9.43, 4.88, 2.85, 26)
text(s, 'The confidence bound\ndetermines b.', 9.43, 5.61, 2.86, 0.74, 16, 'muted')
record(s, 'Choose a greedy arm, then an optimistic challenger', '3 Algorithm', 50, 'Source: Algorithm 1, p.2', 'The two selection rules and matrix update define the learning loop.')
s = new_slide(4, 'One confidence event covers every round', 'The target is a parameter-error bound; its radius comes from noise and regularization.', 45, 'Source: Section 4, Lemma 1, p.5; Eqs. (12)(13)')
rect(s, 0.7, 2.42, 7.05, 2.13, 'cyan_light', radius=True)
text(s, 'With probability ≥ 1−δ, for every t ≤ T:', 0.96, 2.69, 6.5, 0.47, 20, 'cyan', True)
equation(s, '\\|\\widehat\\theta_t-\\theta^{\\star}\\|_{V_{t-1}}\\leq\\frac{\\beta_T}{\\kappa_\\mu}=b', 0.96, 3.45, 6.35, 35)
equation(s, '\\|u\\|_V=\\sqrt{u^{\\top}Vu}', 0.95, 4.97, 6.35, 29)
text(s, 'V sets the shape; b sets the scale.\nδ is the failure probability for all rounds.', 0.96, 5.71, 6.35, 0.86, 20)
rect(s, 7.98, 2.42, 4.66, 4.08, 'white', radius=True)
line(s, 8.37, 5.63, 12.22, 5.63, 'line', 1.2, True)
line(s, 8.65, 6.03, 8.65, 2.9, 'line', 1.2, True)
ellipse = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.72), Inches(3.56), Inches(3.24), Inches(1.57))
ellipse.rotation = 335
ellipse.fill.solid()
ellipse.fill.fore_color.rgb = rgb('cyan_light')
ellipse.line.color.rgb = rgb('cyan')
ellipse.line.width = Pt(2)
dot(s, 10.3, 4.35, 0.065, 'cyan')
dot(s, 10.95, 4.03, 0.067, 'ink')
equation(s, '\\widehat\\theta_t', 9.83, 4.61, 0.8, 24, 'cyan')
equation(s, '\\theta^{\\star}', 11.13, 3.87, 0.8, 24)
text(s, 'Confidence region in parameter space (schematic)', 8.35, 6.05, 3.95, 0.3, 13, 'muted', align=PP_ALIGN.CENTER)
record(s, 'One confidence event covers every round', '4 Confidence Bound', 45, 'Source: Lemma 1, p.5', 'A simultaneous parameter-error bound for all rounds.')
s = new_slide(4, 'Proof sketch: noise plus regularization bias', '', 60, 'Source: Section 4, Lemma 1, pp.5–6; Eqs. (14)(15)')
equation(s, '\\Delta_t=\\widehat\\theta_t-\\theta^{\\star}', 0.69, 1.72, 3.1, 23, 'muted')
equation(s, '\\varepsilon_s=y_s-\\mu(\\langle\\theta^{\\star},\\widetilde{x}_s\\rangle),\\quad\\mathbb{E}[\\varepsilon_s\\mid\\mathcal{F}_{s-1}]=0,\\quad|\\varepsilon_s|\\leq1', 3.92, 1.72, 8.64, 22, 'muted')
rect(s, 0.7, 2.43, 11.94, 1.24, 'white', radius=True)
pill(s, '1  Stationarity + mean value theorem', 0.94, 2.66, 3.45, 'cyan', 'cyan_light')
equation(s, 'H_t\\Delta_t=N_t-\\lambda\\theta^{\\star}', 4.71, 2.68, 7.36, 31)
equation(s, 'N_t:=\\sum_{s<t}\\varepsilon_s\\widetilde{x}_s', 9.38, 2.75, 2.79, 22)
rect(s, 0.7, 3.85, 11.94, 1.39, 'white', radius=True)
pill(s, '2  Curvature lower bound', 0.94, 4.07, 2.42, 'cyan', 'cyan_light')
equation(s, "H_t=\\sum_{s<t}\\mu'(\\xi_{s,t})\\widetilde{x}_s\\widetilde{x}_s^{\\top}+\\lambda I_d\\succeq\\kappa_\\mu V_{t-1}", 4.09, 4.12, 8.07, 28)
rect(s, 0.7, 5.42, 11.94, 1.22, 'cyan_light', radius=True)
pill(s, '3  Cauchy–Schwarz + cancel the error norm', 0.94, 5.62, 4.37, 'cyan', 'white')
equation(s, '\\|\\Delta_t\\|_{V_{t-1}}\\leq\\frac{1}{\\kappa_\\mu}\\left(\\|N_t\\|_{V_{t-1}^{-1}}+\\sqrt{\\lambda\\kappa_\\mu}\\,S\\right)', 5.57, 5.65, 6.56, 29)
record(s, 'Proof sketch: noise plus regularization bias', '4 Confidence Bound', 60, 'Source: Lemma 1, Steps 1–4, pp.5–6', 'Error ≤ inverse curvature × (noise + ridge bias).')
s = new_slide(4, 'The confidence bound determines the UCB bonus', 'Self-normalized concentration controls adaptive noise; Cauchy–Schwarz gives a utility-gap bound.', 45, 'Source: Section 4, Lemma 1, Step 5 and Corollary 1, pp.6–7; Eqs. (12)(16)(17)')
rect(s, 0.7, 2.43, 11.94, 1.9, 'white', radius=True)
text(s, 'Radius numerator = ridge bias + noise bound', 0.96, 2.67, 11.3, 0.39, 20, 'cyan', True)
equation(s, '\\beta_T=\\sqrt{\\lambda\\kappa_\\mu}\\,S+\\sqrt{2\\log(1/\\delta)+d\\log\\left(1+\\frac{4\\kappa_\\mu T}{\\lambda d}\\right)}\\,,\\qquad b=\\frac{\\beta_T}{\\kappa_\\mu}', 0.96, 3.29, 11.26, 31, center=True)
rect(s, 0.7, 4.53, 11.94, 1.39, 'cyan_light', radius=True)
text(s, 'Simultaneously for every pair of arms x, x′:', 0.96, 4.73, 11.25, 0.33, 18, 'cyan', True)
equation(s, "|f(x)-f(x')-\\langle\\widehat\\theta_t,x-x'\\rangle|\\leq b\\,\\|x-x'\\|_{V_{t-1}^{-1}}", 0.96, 5.23, 11.24, 31, center=True)
text(s, 'Hence', 0.98, 6.18, 1.05, 0.34, 19, 'cyan', True)
text(s, 'Estimated utility gap + confidence bonus = the UCB for the second arm.', 2.0, 6.16, 10.25, 0.45, 21)
record(s, 'The confidence bound determines the UCB bonus', '4 Confidence Bound', 45, 'Source: Lemma 1, Step 5 and Corollary 1, pp.6–7', 'The utility-gap confidence bound supplies the exploration bonus.')
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb('navy')
text(s, 'Connecting the four parts', 0.73, 0.87, 11.8, 0.75, 35, 'white', True)
text(s, 'From a preference observation to a comparison backed by a confidence guarantee.', 0.77, 1.81, 11.6, 0.45, 21, 'pale')
for (i, (label, desc)) in enumerate([('Problem', 'Compare two arms;\nobserve a preference.'), ('Assumptions', 'Boundedness;\nlocal curvature.'), ('Algorithm', 'Ridge-MLE estimate;\ngreedy arm + UCB.'), ('Confidence', 'Parameter bounds;\nutility-gap bounds.')]):
    x = 0.76 + 3.14 * i
    text(s, f'0{i + 1}', x, 2.9, 2.65, 0.39, 16, '66C2C1', True)
    text(s, label, x, 3.53, 2.65, 0.61, 30, 'white', True)
    text(s, desc, x, 4.4, 2.73, 0.95, 21, 'pale')
    if i < 3:
        line(s, x + 2.68, 3.84, x + 2.99, 3.84, '66C2C1', 2.1, True)
rect(s, 0.76, 5.78, 11.84, 0.72, '244057', radius=True)
text(s, 'Confidence defines the exploration bonus, which guides the next comparison.', 1.0, 5.98, 11.34, 0.38, 22, 'white', True)
footer(s, len(prs.slides), 0, 'Source: Zhuohua LI · A Clean Proof for Dueling Bandits · §§1–4, pp.1–7', 10, True)
record(s, 'Connecting the four parts', 'Closing', 10, 'Source: Sections 1–4, pp.1–7', 'The model, assumptions, algorithm and confidence bound form one argument.')

def save_notes():
    total = sum((d['seconds'] for d in DECK))
    words = sum((d['spoken_words'] for d in DECK))
    rows = ['# Speaker Notes and Timing — v2', '', 'Scope: A Clean Proof for Dueling Bandits, Sections 1–4.', '', f'{len(DECK)} slides; planned duration: {total // 60}:{total % 60:02d}. The main script contains approximately {words} words. Optional details are separate and are not included in the timing.', '', '| Slide | Topic | Time | Elapsed |', '| --- | --- | --- | --- |']
    elapsed = 0
    for d in DECK:
        elapsed += d['seconds']
        rows.append(f"| {d['page']:02d} | {d['title']} | {d['seconds']} sec | {elapsed // 60}:{elapsed % 60:02d} |")
    rows += ['', 'Section timing: Setup 2:10; Assumptions 2:10; Algorithm 2:30; Confidence 2:30; opening, contents and closing 0:40.', '', 'Slide 6 makes the curvature scope condition explicit. The unconstrained estimator is retained; a bound on the true parameter alone does not establish that scope.', '']
    for d in DECK:
        rows += [f"## {d['page']:02d} | {d['title']}", '', f"Suggested time: {d['seconds']} sec. Key point: {d['key']}", '', d['notes'], '']
        if d['optional']:
            rows += ['**Optional detail (outside the timed script):**', '', d['optional'], '']
        rows += [d['source'], '']
    (ROOT / 'Speaker_Notes_and_Timing.md').write_text('\n'.join(rows), encoding='utf-8')
    (ROOT / 'slide_content.json').write_text(json.dumps(DECK, ensure_ascii=False, indent=2), encoding='utf-8')
    (ASSETS / 'equations.json').write_text(json.dumps(FORMULAS, ensure_ascii=False, indent=2), encoding='utf-8')

def audit_geometry():
    problems = []
    for (i, slide) in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            if shp.left < -Inches(0.01) or shp.top < -Inches(0.01):
                problems.append({'slide': i, 'shape': shp.name, 'issue': 'negative position'})
            if shp.left + shp.width > prs.slide_width + Inches(0.01) or shp.top + shp.height > prs.slide_height + Inches(0.01):
                problems.append({'slide': i, 'shape': shp.name, 'issue': 'outside slide'})
    audit = {'slides': len(prs.slides), 'duration_seconds': sum((d['seconds'] for d in DECK)), 'notes_present': all((bool(sl.notes_slide.notes_text_frame.text.strip()) for sl in prs.slides)), 'smallest_equation_font_pt': min((f['effective_font_pt'] for f in FORMULAS)), 'geometry_issues': problems}
    (PREVIEW / 'structure_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    if problems:
        raise RuntimeError(problems)
    return audit
if __name__ == '__main__':
    for slide in prs.slides:
        for shape in slide.shapes:
            sppr = shape._element.find(qn('p:spPr'))
            if sppr is not None:
                for tag in ('a:effectLst', 'a:effectDag'):
                    existing = sppr.find(qn(tag))
                    if existing is not None:
                        sppr.remove(existing)
                sppr.append(OxmlElement('a:effectLst'))
    prs.save(str(OUTFILE))
    save_notes()
    audit = audit_geometry()
    print(json.dumps({'presentation': str(OUTFILE), 'audit': audit}, ensure_ascii=False, indent=2))
