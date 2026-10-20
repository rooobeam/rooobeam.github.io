# -*- coding: utf-8 -*-
"""Build the first 10-minute presentation, keeping every deliverable under v1.

Python dependencies: python-pptx, matplotlib, numpy, Pillow.
Equations are high-resolution transparent images; all body text and most diagrams
remain native, editable PowerPoint objects. Formula source is kept in this script.
"""

from pathlib import Path
import json
import math
import os
import sys
import tempfile

LOCAL_DEPS = Path(tempfile.gettempdir()) / "codex-dueling-ppt-deps"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import mathtext
from matplotlib.font_manager import FontProperties
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
PREVIEW = ROOT / "preview"
ASSETS.mkdir(exist_ok=True)
PREVIEW.mkdir(exist_ok=True)

OUTFILE = ROOT / "Dueling_Bandits_Sections_1-4_v1.pptx"
W, H = 13.333333, 7.5
TOTAL_SLIDES = 13
FONT = "Microsoft YaHei"
C = {
    "bg": "F7F8FC", "ink": "16263B", "muted": "637187", "line": "DCE3EB",
    "white": "FFFFFF", "teal": "007E87", "teal_light": "E5F2F1",
    "amber": "B57725", "amber_light": "FFF1DD", "blue": "4B62B5",
    "blue_light": "EAEFFA", "cyan": "087EA7", "cyan_light": "E6F3F8",
    "navy": "13273E", "pale": "B4C6D5", "green": "3D875C",
}
SECTION = {
    1: ("PROBLEM SETUP", "问题设定", "teal", "teal_light"),
    2: ("ASSUMPTIONS", "假设", "amber", "amber_light"),
    3: ("ALGORITHM", "算法", "blue", "blue_light"),
    4: ("CONFIDENCE BOUND", "置信界", "cyan", "cyan_light"),
}

matplotlib.rcParams.update({
    "mathtext.fontset": "stix", "font.family": "DejaVu Sans",
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.labelcolor": "#637187", "xtick.color": "#637187", "ytick.color": "#637187",
    "axes.edgecolor": "#BCC8D2", "figure.facecolor": "none", "savefig.facecolor": "none",
})

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "Contextual Dueling Bandits：问题、假设、算法与置信界"
prs.core_properties.subject = "A Clean Proof for Dueling Bandits，第 1–4 节，10 分钟讲解"
prs.core_properties.author = ""
prs.core_properties.keywords = "Dueling Bandits, Ridge-MLE, UCB, Confidence Bound"
prs.core_properties.comments = "v1：四部分篇幅相对均衡；逐页讲稿位于备注。"

DECK = []
FORMULAS = []


def rgb(value):
    value = C.get(value, value)
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill="white", line=None, radius=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if radius:
        shp.adjustments[0] = 0.10
    return shp


def text(slide, value, x, y, w, h, size=22, color="ink", bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, font=FONT):
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(4)
        p.line_spacing = 1.12
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        rpr = run._r.get_or_add_rPr()
        ea = OxmlElement("a:ea")
        ea.set("typeface", font)
        rpr.append(ea)
    return shp


def line(slide, x1, y1, x2, y2, color="line", width=1.4, arrow=False):
    shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "med")
        end.set("len", "med")
        shp.line._get_or_add_ln().append(end)
    return shp


def dot(slide, x, y, r=0.06, color="teal"):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r),
                                 Inches(2*r), Inches(2*r))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    return shp


def equation(slide, latex, x, y, maxw, size=26, color="ink", maxh=None, center=False):
    """Render at the requested point size, shrinking only when necessary."""
    # Older mathtext versions render \| as a single bar. Use an explicit norm.
    latex = latex.replace(r"\|", r"\Vert ")
    idx = len(FORMULAS) + 1
    path = ASSETS / f"equation-{idx:03d}.png"
    with matplotlib.rc_context({"text.color": "#" + C.get(color, color)}):
        mathtext.math_to_image("$" + latex + "$", str(path),
                              prop=FontProperties(size=size), dpi=240)
    im = Image.open(path).convert("RGBA")
    padded = Image.new("RGBA", (im.width+18, im.height+18))
    padded.alpha_composite(im, (9, 9))
    padded.save(path)
    w, h = padded.width/240, padded.height/240
    scale = min(1.0, maxw/w, maxh/h if maxh else 1.0)
    w, h = w*scale, h*scale
    actual_x = x + (maxw-w)/2 if center else x
    pic = slide.shapes.add_picture(str(path), Inches(actual_x), Inches(y),
                                   width=Inches(w), height=Inches(h))
    pic.name = f"Equation {idx}: {latex}"
    FORMULAS.append({"file": path.name, "latex": latex, "slide": len(prs.slides),
                     "effective_font_pt": round(size*scale, 1)})
    return pic


def pill(slide, label, x, y, w, accent="teal", light="teal_light", size=13):
    rect(slide, x, y, w, .34, light, radius=True)
    text(slide, label, x+.09, y+.045, w-.18, .25, size=size, color=accent, bold=True)


def footer(slide, number, section, source, seconds, dark=False):
    ink = "pale" if dark else "muted"
    line(slide, .64, 7.02, 12.69, 7.02, "426077" if dark else "line", .8)
    text(slide, source, .67, 7.13, 9.9, .2, size=10, color=ink)
    text(slide, f"{seconds} 秒   ·   {number:02d} / {TOTAL_SLIDES}", 11.0, 7.10, 1.65, .25,
         size=11, color=ink, align=PP_ALIGN.RIGHT)
    for i in range(1, 5):
        col = SECTION[i][2] if i == section else ("426077" if dark else "DCE3EB")
        rect(slide, .67+(i-1)*.53, 6.88, .40, .045, col)


def new_slide(section, title, subtitle, seconds, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("bg")
    english, chinese, accent, light = SECTION[section]
    text(slide, f"{section:02d}  {english}  /  {chinese}", .68, .38, 11.9, .30,
         size=12, color=accent, bold=True)
    text(slide, title, .66, .94, 12.05, .66, size=31, bold=True)
    text(slide, subtitle, .69, 1.67, 11.95, .49, size=16, color="muted")
    footer(slide, len(prs.slides), section, source, seconds)
    return slide


def record(slide, title, section, seconds, notes, source, key):
    slide.notes_slide.notes_text_frame.text = (
        f"第 {len(prs.slides)} 页｜{title}\n建议用时：{seconds} 秒\n"
        f"核心信息：{key}\n\n{notes}\n\n来源：{source}"
    )
    DECK.append({"page": len(prs.slides), "title": title, "section": section,
                 "seconds": seconds, "key": key, "notes": notes, "source": source})


def plot_sigmoid():
    fig, ax = plt.subplots(figsize=(4.8, 2.15), dpi=200)
    z = np.linspace(-5, 5, 350)
    ax.plot(z, 1/(1+np.exp(-z)), color="#"+C["teal"], lw=3)
    ax.axhline(.5, color="#C4CED6", ls="--", lw=1)
    ax.axvline(0, color="#C4CED6", ls="--", lw=1)
    ax.scatter([0], [.5], color="#"+C["teal"], s=36, zorder=4)
    ax.set(xlim=(-5, 5), ylim=(-.02, 1.03), xticks=[-4, 0, 4], yticks=[0, .5, 1])
    ax.tick_params(labelsize=12)
    ax.set_xlabel(r"utility gap $z$", fontsize=12)
    ax.set_ylabel(r"$\mu(z)$", fontsize=14, rotation=0, labelpad=17)
    fig.tight_layout(pad=.7)
    target = ASSETS / "logistic-link.png"
    fig.savefig(target, transparent=True)
    plt.close(fig)
    return target


def plot_curvature():
    fig, ax = plt.subplots(figsize=(5.0, 2.6), dpi=220)
    z = np.linspace(-6, 6, 400)
    p = 1/(1+np.exp(-z))
    k = np.exp(-2)/(1+np.exp(-2))**2
    ax.axvspan(-2, 2, color="#F7E7CC", alpha=.75)
    ax.plot(z, p*(1-p), color="#"+C["amber"], lw=3)
    ax.hlines(k, -2, 2, color="#"+C["ink"], ls="--", lw=1.4)
    ax.vlines([-2, 2], 0, k, color="#BBA27F", ls=":", lw=1.2)
    ax.annotate(r"$\kappa_\mu$", (2.1, k+.012), fontsize=18, color="#"+C["ink"])
    ax.set(xlim=(-6, 6), ylim=(0, .28), xticks=[-6, -2, 0, 2, 6], yticks=[0, .1, .2])
    ax.tick_params(labelsize=12)
    ax.set_xlabel(r"$z$  (illustration: $S=1$)", fontsize=12)
    ax.set_ylabel(r"$\mu'(z)$", fontsize=16, rotation=0, labelpad=15)
    fig.tight_layout(pad=.7)
    target = ASSETS / "local-curvature.png"
    fig.savefig(target, transparent=True)
    plt.close(fig)
    return target


# 01 — opening, 10 sec
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb("navy")
rect(s, .68, .70, .10, 3.10, "teal")
text(s, "A CLEAN PROOF FOR DUELING BANDITS", .98, .72, 11.3, .35,
     size=15, color="pale", bold=True)
text(s, "从成对偏好\n到置信上界", .95, 1.42, 9.4, 1.9,
     size=44, color="white", bold=True)
text(s, "Contextual Dueling Bandits", 1.0, 3.48, 10.7, .48,
     size=27, color="pale", font="Calibri")
text(s, "第 1–4 节  ·  约 10 分钟  ·  v1", 1.0, 4.12, 10.5, .36,
     size=17, color="pale")
for i, (cn, en, timing) in enumerate([
    ("问题设定", "What do we observe?", "2 分 10 秒"),
    ("假设", "What makes it learnable?", "2 分 10 秒"),
    ("算法", "How do we choose?", "2 分 30 秒"),
    ("置信界", "Why trust the estimate?", "2 分 30 秒"),
]):
    x = .98 + i*3.08
    text(s, f"0{i+1}", x, 5.12, .55, .35, size=15, color="66C2C1", bold=True)
    text(s, cn, x, 5.60, 2.7, .48, size=24, color="white", bold=True)
    text(s, timing, x, 6.20, 2.7, .30, size=13, color="pale")
footer(s, 1, 0, "Zhuohua LI · A Clean Proof for Dueling Bandits · February 2026 · §§1–4", 10, True)
record(s, "从成对偏好到置信上界", "开场", 10,
       "今天用十分钟介绍文稿的前四节，主线是从成对偏好出发，走到有置信保证的选臂策略。",
       "原稿 §§1–4，pp.1–7", "偏好反馈 → 参数估计 → 带置信保证的比较")


# 02 — contents, 20 sec
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb("bg")
text(s, "CONTENTS  /  讲解路线", .68, .38, 11.9, .30, 12, "teal", True)
text(s, "目录", .66, .94, 12.05, .66, 34, bold=True)
text(s, "四个问题，串起从偏好反馈到置信上界的逻辑。", .69, 1.67, 11.95, .49, 18, "muted")
agenda = [
    ("Problem Setup: Contextual Dueling Bandits", "观察什么：候选特征、成对偏好与遗憾", "03–04 页", "2 分 10 秒"),
    ("Assumptions", "依赖什么：有界性与局部曲率下界", "05–06 页", "2 分 10 秒"),
    ("Algorithm: Ridge-MLE Dueling UCB", "怎样选择：参数估计、贪心第一臂与 UCB 第二臂", "07–09 页", "2 分 30 秒"),
    ("Confidence Bound", "为何可信：参数置信椭球与效用差置信界", "10–12 页", "2 分 30 秒"),
]
for i, (heading, description, pages, timing) in enumerate(agenda, start=1):
    y = 2.35 + (i-1)*1.06
    accent = SECTION[i][2]
    rect(s, .70, y, 11.94, .92, "white", radius=True)
    rect(s, .70, y+.15, .045, .62, accent)
    text(s, f"0{i}", .96, y+.23, .72, .49, 27, accent, True)
    text(s, heading, 1.89, y+.12, 8.48, .37, 21, bold=True, font="Calibri")
    text(s, description, 1.91, y+.56, 8.40, .29, 15, "muted")
    text(s, pages, 10.62, y+.19, 1.68, .32, 17, accent, True, PP_ALIGN.RIGHT)
    text(s, timing, 10.62, y+.59, 1.68, .25, 12, "muted", align=PP_ALIGN.RIGHT)
footer(s, len(prs.slides), 0, "A Clean Proof for Dueling Bandits · §§1–4 · 目录与页码", 20)
record(s, "目录", "目录", 20,
       "讲解分成四部分：先说明能观察什么、希望优化什么；再看有界性和曲率假设；接着介绍参数估计和两个候选的选择；最后用置信界解释探索项为什么成立。四部分各用两分钟左右。",
       "原稿 §§1–4，pp.1–7", "问题设定 → 假设 → 算法 → 置信界")


# 03 — setup 1/2, 65 sec
s = new_slide(1, "每轮比较两个候选，只收到一个偏好标签",
              "Contextual：候选带有特征，决策集合也可以随轮次变化。", 65,
              "原稿 §1，p.1；式 (1)–(3)")
xs = [.70, 4.97, 9.24]
widths = [3.40, 3.40, 3.40]
for x, w in zip(xs, widths):
    rect(s, x, 2.43, w, 2.73, "white", radius=True)
for i, label in enumerate(["观察候选集合", "选择一对候选", "观察偏好结果"]):
    pill(s, f"{i+1}  {label}", xs[i]+.22, 2.66, 2.65)
equation(s, r"\mathcal{X}_t\subseteq\mathbb{R}^{d}", .96, 3.30, 2.95, 31)
text(s, "每个候选由特征 x 表示", .96, 4.17, 2.98, .75, 18)
equation(s, r"x_{t,1},\ x_{t,2}\in\mathcal{X}_t", 5.21, 3.30, 2.98, 29)
text(s, "例如：两条待比较的推荐", 5.23, 4.17, 2.98, .75, 18)
equation(s, r"y_t\in\{0,1\}", 9.50, 3.30, 2.95, 31)
text(s, "y = 1：第一个候选胜出", 9.48, 4.17, 3.0, .75, 18)
line(s, 4.24, 3.72, 4.80, 3.72, "teal", 2.4, True)
line(s, 8.50, 3.72, 9.07, 3.72, "teal", 2.4, True)
rect(s, .70, 5.47, 11.94, 1.06, "teal_light", radius=True)
text(s, "不可直接观测的效用", .94, 5.68, 2.75, .44, 18, "teal", True)
equation(s, r"f(x)=\langle\theta^{\star},x\rangle", 3.62, 5.67, 3.27, 27)
text(s, "比较使用差分特征", 7.02, 5.70, 2.43, .40, 16, "teal")
equation(s, r"\widetilde{x}_t=x_{t,1}-x_{t,2}", 9.10, 5.71, 3.22, 25)
record(s, "每轮比较两个候选，只收到一个偏好标签", "1 Problem Setup", 65,
       "先看一轮交互。学习者在第 t 轮拿到候选集合，每个候选是一个 d 维特征向量。这里的 contextual 指决策依赖这些特征，而且每一轮的候选集合可以变化。学习者选择两个候选，环境只返回一个零一标签：一表示第一个胜出，零表示第二个胜出。\n\n真实的潜在效用假设为线性函数，未知的是参数 θ⋆。我们没有直接观察效用数值，而是利用比较结果推断它。由于两个效用相减等于参数与特征差的内积，后面算法处理的输入自然是差分特征 x̃。候选、选臂以及可能的算法随机性应当纳入观察本轮标签之前的信息，以满足后面条件期望和鞅分析的要求。",
       "原稿 §1，p.1", "一轮交互的可观测量与潜在效用必须区分")


# 04 — setup 2/2, 65 sec
s = new_slide(1, "用效用差建模偏好，用遗憾衡量选择",
              "胜率描述反馈机制；遗憾描述我们最终希望优化的目标。", 65,
              "原稿 §1，p.1；BTL 模型与 regret 定义")
rect(s, .70, 2.40, 5.77, 4.11, "white", radius=True)
rect(s, 6.68, 2.40, 5.96, 4.11, "white", radius=True)
pill(s, "LOGISTIC / BTL 反馈", .94, 2.62, 2.73)
equation(s, r"\Pr(y_t=1\mid\mathcal{F}_{t-1})=\mu(\langle\theta^{\star},\widetilde{x}_t\rangle)",
         .94, 3.16, 5.22, 25)
equation(s, r"\mu(z)=\frac{1}{1+e^{-z}}", .96, 3.73, 4.95, 25)
sigmoid = plot_sigmoid()
s.shapes.add_picture(str(sigmoid), Inches(1.02), Inches(4.30), width=Inches(4.96), height=Inches(1.78))
text(s, "效用相同：胜率 1/2；效用越高：越可能胜出。", .98, 6.13, 5.19, .25, 12, "muted")
pill(s, "优化目标：两个候选都接近最优", 6.94, 2.62, 4.66)
equation(s, r"x_t^{\star}\in\arg\max_{x\in\mathcal{X}_t}f(x)", 6.96, 3.20, 5.10, 27)
text(s, "单步遗憾 = 两个候选各自的效用损失", 6.96, 4.01, 5.20, .37, 18)
equation(s, r"r_t=[f(x_t^{\star})-f(x_{t,1})]", 6.97, 4.51, 5.10, 27)
equation(s, r"\quad+[f(x_t^{\star})-f(x_{t,2})]", 6.97, 4.99, 5.10, 27)
equation(s, r"R_T=\sum_{t=1}^{T}r_t", 6.97, 5.55, 4.90, 25)
record(s, "用效用差建模偏好，用遗憾衡量选择", "1 Problem Setup", 65,
       "有了效用差，还需要说明它怎样产生偏好。文稿采用 Logistic，也就是 BTL 模型：把两个候选的效用差输入 S 形函数，得到第一个胜出的概率。差为零时，双方各有一半胜率；差越大，第一个越可能胜出，但并不保证每次都赢。式中的历史信息要包含本轮反馈之前已经确定的候选与选臂。\n\n优化目标用右边的遗憾表示。每一轮都有一个真实效用最高的候选，单步遗憾是所选两个候选相对它的效用损失之和，累计遗憾就是跨轮次相加。这意味着两个候选都选得好才是目标，不能只保证其中一个好。这里沿用原稿的未除以二的遗憾定义，并默认涉及的最大值可以取到；若集合无限，可由紧性等条件保证。",
       "原稿 §1，p.1", "偏好是随机反馈；遗憾是潜在效用差")


# 05 — assumptions 1/2, 65 sec
s = new_slide(2, "有界性控制尺度：先约束特征与真实参数",
              "这些界会分别进入曲率控制、设计矩阵上界和正则化偏差。", 65,
              "原稿 §1 式 (1)(2) 与 §2 Assumption 1，pp.1–2")
rect(s, .70, 2.43, 5.15, 3.33, "white", radius=True)
pill(s, "给定的尺度约束", .96, 2.68, 2.31, "amber", "amber_light")
equation(s, r"\|x\|_2\leq 1", .96, 3.36, 4.30, 34)
text(s, "候选特征有界", 3.55, 3.49, 1.95, .4, 18, "muted")
equation(s, r"\|\theta^{\star}\|_2\leq S", .96, 4.40, 4.30, 34)
text(s, "S 是已知常数", 3.55, 4.53, 1.95, .4, 18, "muted")
line(s, 6.02, 4.07, 6.56, 4.07, "amber", 2.6, True)
rect(s, 6.77, 2.43, 5.87, 3.33, "amber_light", radius=True)
text(s, "推得差分特征与真实得分有界", 7.03, 2.71, 5.17, .45, 21, "amber", True)
equation(s, r"\|\widetilde{x}_t\|_2\leq\|x_{t,1}\|_2+\|x_{t,2}\|_2\leq2", 7.03, 3.48, 5.22, 27)
equation(s, r"|\langle\theta^{\star},\widetilde{x}_t\rangle|\leq2S", 7.03, 4.45, 5.18, 31)
text(s, "直观作用", .94, 6.07, 1.55, .35, 18, "amber", True)
text(s, "限定真实效用的范围，并为后续的误差控制提供尺度。", 2.52, 6.07, 9.76, .40, 21)
record(s, "有界性控制尺度：先约束特征与真实参数", "2 Assumptions", 65,
       "先看最基础的尺度条件。每个候选的特征范数不超过一，真实参数范数不超过已知常数 S。由三角不等式，两个候选的差分特征范数不超过二；再用柯西施瓦茨不等式，真实效用差的绝对值就不超过二 S。\n\n这些条件并非只是记号上的方便。差分特征有界，后面才能控制设计矩阵增长以及行列式；真实参数有界，后面才能控制岭正则项带来的偏差；真实得分落在有限区间内，也让 Logistic 曲线在这个区间上拥有正的最小导数。注意这里约束的是 θ⋆，并没有自动约束算法求出的 θ̂。这个区别正好引出下一页的曲率条件。",
       "原稿 §1–2，pp.1–2", "有界性同时支撑特征、曲率和正则化项的控制")


# 06 — assumptions 2/2, 65 sec
s = new_slide(2, "曲率下界把预测变化连接到参数误差",
              "Logistic 曲线在两端变平；证明需要在相关区间内保留足够的曲率。", 65,
              "原稿 §2 Assumption 2 与 §4 Lemma 1 Step 2，pp.2、5")
rect(s, .70, 2.43, 5.66, 2.93, "white", radius=True)
curve = plot_curvature()
s.shapes.add_picture(str(curve), Inches(.88), Inches(2.60), width=Inches(5.25), height=Inches(2.52))
rect(s, 6.57, 2.43, 6.07, 2.93, "white", radius=True)
pill(s, "ASSUMPTION 2", 6.83, 2.66, 2.35, "amber", "amber_light")
equation(s, r"\mu'(z)=\mu(z)(1-\mu(z))\geq\kappa_\mu>0", 6.83, 3.26, 5.43, 28)
text(s, "证明实际用到的中间点：", 6.84, 4.03, 5.18, .42, 20)
equation(s, r"\xi_{s,t}\ \mathrm{between}\ \langle\theta^{\star},\widetilde{x}_s\rangle\ \mathrm{and}\ \langle\widehat\theta_t,\widetilde{x}_s\rangle", 6.84, 4.63, 5.32, 24)
rect(s, .70, 5.59, 11.94, 1.04, "amber_light", radius=True)
text(s, "适用范围", .94, 5.80, 1.35, .35, 17, "amber", True)
text(s, "仅有 θ⋆ 有界还不够：也需保证上述中间点处的曲率下界。\n以下置信界以这一范围条件成立为前提。", 2.39, 5.76, 9.86, .72, 17)
record(s, "曲率下界把预测变化连接到参数误差", "2 Assumptions", 65,
       "Logistic 函数的导数在零附近较大，到两端逐渐接近零。如果曲线太平，参数变动很多，预测概率却只变化一点，就难以从预测误差反推出参数误差。因此假设在证明涉及的得分区间内，导数至少为一个正数 κμ。这个条件会把后面二次型的下界变成 κμ 乘以参数误差的平方范数。\n\n这里需要留意原稿的一处范围问题。真实参数有界只保证真实得分位于正负二 S 之间，但原算法的估计是在整个参数空间上优化，未保证 θ̂ 也在半径 S 的球内。均值定理用到的是两种得分之间的中间点，因此曲率下界必须覆盖这些点。我们以下直接把这个范围要求作为假设，不把真实参数有界当成它的充分证明。图中 S 等于一只是曲线示意，不是算法参数设定。",
       "原稿 §2 p.2；§4 p.5。范围注记为讲解时的严谨性补充。", "κμ 必须适用于估计得分与真实得分之间的中间点")


# 07 — algorithm 1/3, 50 sec
s = new_slide(3, "把偏好标签转成带正则项的逻辑回归",
              "用前 t−1 轮数据，估计第 t 轮使用的参数。", 50,
              "原稿 §3.1，pp.2–4；式 (4)–(8)")
rect(s, .70, 2.44, 11.94, 1.38, "white", radius=True)
pill(s, "1  单次比较的负对数似然", .93, 2.65, 3.55, "blue", "blue_light")
equation(s, r"\ell(y,z)=-y\log\mu(z)-(1-y)\log(1-\mu(z))", 4.74, 2.80, 7.49, 27)
text(s, "预测得分", .94, 3.18, 1.21, .28, 14, "muted")
equation(s, r"z=\langle\theta,\widetilde{x}_s\rangle", 2.12, 3.16, 2.20, 23)
rect(s, .70, 4.01, 11.94, 1.59, "blue_light", radius=True)
pill(s, "2  累加损失，再加入岭正则", .93, 4.24, 3.72, "blue", "white")
equation(s, r"\widehat\theta_t=\arg\min_{\theta\in\mathbb{R}^{d}}\left\{\sum_{s<t}\ell(y_s,\langle\theta,\widetilde{x}_s\rangle)+\frac{\lambda}{2}\|\theta\|_2^2\right\}",
         4.82, 4.38, 7.44, 30)
text(s, "λ > 0", .95, 6.01, 1.31, .43, 22, "blue", True)
text(s, "使目标严格凸、稳定参数估计；同时在置信界中产生一项偏差。", 2.28, 6.02, 10.03, .52, 20)
record(s, "把偏好标签转成带正则项的逻辑回归", "3 Algorithm", 50,
       "现在开始构造估计量。给定一组参数，每次偏好标签都是伯努利观测，成功概率是该参数预测的 Logistic 概率。最大化历史数据的条件似然，等价于最小化这里的负对数似然，也就是熟悉的二元交叉熵。历史比较是自适应选择的，因此这里用条件概率的链式分解，不需要声称这些观测无条件独立。\n\n将前 t 减一轮的损失相加，再加入二分之 λ 乘参数平方范数，就得到 Ridge-MLE。λ 为正使目标严格凸，最小点唯一且稳定。正则化也有代价：它把估计向零拉，置信半径必须把这种偏差算进去。第一轮没有数据，估计就是零。",
       "原稿 §3.1，pp.2–4", "伯努利偏好的条件似然产生逻辑损失，岭项稳定估计")


# 08 — algorithm 2/3, 50 sec
s = new_slide(3, "一阶条件是置信界证明的起点",
              "Logistic 的特殊结构，让损失导数简化为“预测概率 − 实际标签”。", 50,
              "原稿 §3.2，p.4；式 (9)(10)")
for y, tag in [(2.44, "对得分求导"), (3.66, "对参数求梯度"), (5.04, "最小点满足")]:
    rect(s, .70, y, 11.94, 1.07 if y < 5 else 1.37,
         "white" if y < 5 else "blue_light", radius=True)
    text(s, tag, .95, y+.31, 2.68, .43, 20, "blue", True)
equation(s, r"\frac{\partial\ell(y,z)}{\partial z}=\mu(z)-y", 4.15, 2.65, 7.86, 30)
equation(s, r"\nabla_\theta\ell(y_s,\langle\theta,\widetilde{x}_s\rangle)=[\mu(\langle\theta,\widetilde{x}_s\rangle)-y_s]\widetilde{x}_s", 4.15, 3.94, 7.93, 27)
equation(s, r"\sum_{s<t}[\mu(\langle\widehat\theta_t,\widetilde{x}_s\rangle)-y_s]\widetilde{x}_s+\lambda\widehat\theta_t=0", 4.15, 5.35, 7.92, 30)
record(s, "一阶条件是置信界证明的起点", "3 Algorithm", 50,
       "这一页只保留求导中最重要的三步。第一步，对逻辑损失的得分求导，因为 Logistic 的导数等于 μ 乘一减 μ，分母消掉之后，结果就是预测概率减实际标签。第二步，利用链式法则，对参数的梯度再乘上差分特征。第三步，把所有历史样本和正则项的梯度相加，在最小点处等于零。\n\n这条等式后面会被改写：把真实概率加进来再减掉，就能把估计偏差放到左边，把观测噪声和正则化偏差放到右边。因此它既是优化问题的最优性条件，也是统计误差证明的入口。此处梯度等于零对应原稿的无约束最小化；若以后改成球约束估计，证明也需相应改用最优性不等式。",
       "原稿 §3.2，p.4", "掌握最终驻点等式，不需要逐行展开所有求导")


# 09 — algorithm 3/3, 50 sec
s = new_slide(3, "先选当前最优，再选值得比较的挑战者",
              "Ridge-MLE Dueling UCB：第一臂利用当前估计，第二臂借助不确定性探索。", 50,
              "原稿 §3 Algorithm 1，p.2；b = βT / κμ")
rect(s, .70, 2.42, 8.26, 1.37, "white", radius=True)
pill(s, "第一臂：预测效用最高", .93, 2.62, 3.23, "blue", "blue_light")
equation(s, r"x_{t,1}\in\arg\max_{x\in\mathcal{X}_t}\langle\widehat\theta_t,x\rangle", 4.42, 2.81, 4.15, 28)
rect(s, .70, 3.99, 8.26, 2.49, "blue_light", radius=True)
pill(s, "第二臂：相对第一臂的 UCB 最大", .93, 4.21, 4.54, "blue", "white")
equation(s, r"x_{t,2}\in\arg\max_{x\in\mathcal{X}_t}\{\langle\widehat\theta_t,x-x_{t,1}\rangle", .95, 4.94, 7.61, 29)
equation(s, r"\qquad\qquad\quad+b\,\|x-x_{t,1}\|_{V_{t-1}^{-1}}\}", .95, 5.67, 7.61, 29)
rect(s, 9.18, 2.42, 3.46, 4.06, "white", radius=True)
text(s, "每轮维护的信息", 9.43, 2.67, 2.91, .43, 21, "blue", True)
equation(s, r"V_0=\frac{\lambda}{\kappa_\mu}I_d", 9.43, 3.38, 2.82, 28)
text(s, "收到偏好后更新", 9.43, 4.30, 2.86, .38, 18)
equation(s, r"V_t=V_{t-1}+\widetilde{x}_t\widetilde{x}_t^{\top}", 9.43, 4.88, 2.85, 26)
text(s, "b 的数值由下一部分\n的置信界确定。", 9.43, 5.61, 2.86, .74, 16, "muted")
record(s, "先选当前最优，再选值得比较的挑战者", "3 Algorithm", 50,
       "有了估计参数，第一臂直接选当前预测效用最高的候选。第二臂则相对于第一臂做乐观比较：预测的效用差，再加一个不确定性奖励，选这个和最大的候选。由于第一臂已经是预测最优，其他候选的预测差通常不为正；挑战者仍可能被选中，是因为它的相对效用还不够确定。\n\n不确定性由差分方向在 V 逆矩阵下的长度来衡量。V 从 λ 除以 κμ 倍的单位阵开始，每次增加实际比较方向的外积。某类差分方向被观察得更多，在该方向上的不确定性会下降。收到标签后更新数据和 V，进入下一轮重新估计。剩下的问题是：前面的系数 b 应该多大，才能使这个上界可信？",
       "原稿 Algorithm 1，p.2", "两种选臂规则与设计矩阵更新构成完整循环")


# 10 — confidence 1/3, 45 sec
s = new_slide(4, "置信椭球：一次高概率事件，覆盖所有轮次",
              "先明确要证明什么，再看误差如何被噪声与正则化项控制。", 45,
              "原稿 §4 Lemma 1，p.5；式 (12)(13)")
rect(s, .70, 2.42, 7.05, 2.13, "cyan_light", radius=True)
text(s, "以至少 1−δ 的概率，对所有 t ≤ T 同时成立：", .96, 2.69, 6.50, .47, 20, "cyan", True)
equation(s, r"\|\widehat\theta_t-\theta^{\star}\|_{V_{t-1}}\leq\frac{\beta_T}{\kappa_\mu}=b", .96, 3.45, 6.35, 35)
equation(s, r"\|u\|_V=\sqrt{u^{\top}Vu}", .95, 4.97, 6.35, 29)
text(s, "V 决定椭球形状，b 决定整体尺度。\nδ 是失败概率；保证覆盖整个时间段。", .96, 5.71, 6.35, .86, 20)
rect(s, 7.98, 2.42, 4.66, 4.08, "white", radius=True)
line(s, 8.37, 5.63, 12.22, 5.63, "line", 1.2, True)
line(s, 8.65, 6.03, 8.65, 2.90, "line", 1.2, True)
ellipse = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.72), Inches(3.56), Inches(3.24), Inches(1.57))
ellipse.rotation = 335
ellipse.fill.solid()
ellipse.fill.fore_color.rgb = rgb("cyan_light")
ellipse.line.color.rgb = rgb("cyan")
ellipse.line.width = Pt(2)
dot(s, 10.30, 4.35, .065, "cyan")
dot(s, 10.95, 4.03, .067, "ink")
equation(s, r"\widehat\theta_t", 9.83, 4.61, .80, 24, "cyan")
equation(s, r"\theta^{\star}", 11.13, 3.87, .80, 24)
text(s, "参数空间中的置信区域（示意）", 8.35, 6.05, 3.95, .30, 13, "muted", align=PP_ALIGN.CENTER)
record(s, "置信椭球：一次高概率事件，覆盖所有轮次", "4 Confidence Bound", 45,
       "置信界要证明的是：以至少一减 δ 的概率，所有轮次的估计参数与真实参数之间，都满足同一个椭球范数误差界。δ 表示失败概率，不能把它与置信水平混淆。右图是在参数空间里的示意：以估计参数为中心，希望真实参数落在置信椭球内。\n\n这里不是普通欧氏距离，而是 V 加权的距离。V 记录哪些差分方向已经被充分观察，因此它决定椭球的形状，b 决定整体大小。这个结论是对所有轮次同时成立的；后面使用时间一致的自归一化不等式，不需要简单地给每轮另外分配一次失败概率。接着看这个半径由什么构成。",
       "原稿 Lemma 1，p.5", "同时覆盖全部轮次的参数误差界")


# 11 — confidence 2/3, 60 sec
s = new_slide(4, "证明骨架：先把误差拆成噪声与正则偏差",
              "", 60,
              "原稿 §4 Lemma 1，pp.5–6；式 (14)(15)")
equation(s, r"\Delta_t=\widehat\theta_t-\theta^{\star}", .69, 1.72, 3.10, 23, "muted")
equation(s, r"\varepsilon_s=y_s-\mu(\langle\theta^{\star},\widetilde{x}_s\rangle),\quad\mathbb{E}[\varepsilon_s\mid\mathcal{F}_{s-1}]=0,\quad|\varepsilon_s|\leq1", 3.92, 1.72, 8.64, 22, "muted")
rect(s, .70, 2.43, 11.94, 1.24, "white", radius=True)
pill(s, "1  驻点等式 + 均值定理", .94, 2.66, 3.45, "cyan", "cyan_light")
equation(s, r"H_t\Delta_t=N_t-\lambda\theta^{\star}", 4.71, 2.68, 7.36, 31)
equation(s, r"N_t:=\sum_{s<t}\varepsilon_s\widetilde{x}_s", 9.38, 2.75, 2.79, 22)
rect(s, .70, 3.85, 11.94, 1.39, "white", radius=True)
pill(s, "2  曲率给下界", .94, 4.07, 2.42, "cyan", "cyan_light")
equation(s, r"H_t=\sum_{s<t}\mu'(\xi_{s,t})\widetilde{x}_s\widetilde{x}_s^{\top}+\lambda I_d\succeq\kappa_\mu V_{t-1}",
         4.09, 4.12, 8.07, 28)
rect(s, .70, 5.42, 11.94, 1.22, "cyan_light", radius=True)
pill(s, "3  柯西施瓦茨 + 消去误差范数", .94, 5.62, 4.37, "cyan", "white")
equation(s, r"\|\Delta_t\|_{V_{t-1}}\leq\frac{1}{\kappa_\mu}\left(\|N_t\|_{V_{t-1}^{-1}}+\sqrt{\lambda\kappa_\mu}\,S\right)",
         5.57, 5.65, 6.56, 29)
record(s, "证明骨架：先把误差拆成噪声与正则偏差", "4 Confidence Bound", 60,
       "证明可以压成三步。先定义估计误差 Δ，以及观测减真实概率得到的噪声 ε。这个噪声的条件均值为零、绝对值不超过一，所以可以使用条件次高斯的鞅集中工具。第一步，从刚才的驻点等式出发，加上再减去真实概率，并使用均值定理，得到 H 乘 Δ 等于噪声累积 N 减去 λθ⋆。\n\n第二步，在前面说明的曲率范围条件下，H 至少是 κμ 倍的 V。第三步，两侧左乘 Δ 的转置。左侧至少为 κμ 乘误差范数的平方；右侧用加权柯西施瓦茨，得到误差范数乘以噪声范数和正则偏差之和。非零时消去误差范数，零时结论显然成立。又因为 V 至少是 λ 除以 κμ 倍的单位阵，正则项最多为根号 λκμ 乘 S。这样只剩噪声项还需要概率上界。",
       "原稿 Lemma 1 Steps 1–4，pp.5–6", "最终误差 ≤ 曲率倒数 ×（随机噪声 + 正则化偏差）")


# 12 — confidence 3/3, 45 sec
s = new_slide(4, "噪声有了上界，UCB 的探索项就确定了",
              "自归一化不等式控制自适应数据的噪声；加权柯西施瓦茨把参数界转成效用差界。", 45,
              "原稿 §4 Lemma 1 Step 5 与 Corollary 1，pp.6–7；式 (12)(16)(17)")
rect(s, .70, 2.43, 11.94, 1.90, "white", radius=True)
text(s, "置信半径 = 正则化偏差 + 噪声上界", .96, 2.67, 11.30, .39, 20, "cyan", True)
equation(s, r"\beta_T=\sqrt{\lambda\kappa_\mu}\,S+\sqrt{2\log(1/\delta)+d\log\left(1+\frac{4\kappa_\mu T}{\lambda d}\right)}\,,\qquad b=\frac{\beta_T}{\kappa_\mu}",
         .96, 3.29, 11.26, 31, center=True)
rect(s, .70, 4.53, 11.94, 1.39, "cyan_light", radius=True)
text(s, "对任意候选对 x, x′，同时成立：", .96, 4.73, 11.25, .33, 18, "cyan", True)
equation(s, r"|f(x)-f(x')-\langle\widehat\theta_t,x-x'\rangle|\leq b\,\|x-x'\|_{V_{t-1}^{-1}}", .96, 5.23, 11.24, 31, center=True)
text(s, "因此", .98, 6.18, 1.05, .34, 19, "cyan", True)
text(s, "预测的效用差 + 置信半径，就是第二个候选使用的 UCB。", 2.0, 6.16, 10.25, .45, 21)
record(s, "噪声有了上界，UCB 的探索项就确定了", "4 Confidence Bound", 45,
       "最后用自归一化不等式控制 N 的 V 逆范数。它能处理比较方向依赖历史数据的情形，给出同时覆盖所有时间的噪声上界。再用差分特征范数不超过二，以及行列式与迹的关系，把结果化成这里显式的对数项。它与正则化项相加就是 β，再除以 κμ 得到 b。\n\n参数置信界通过一次加权柯西施瓦茨，就变成下面的效用差置信界：真实效用差与预测差的偏离，不超过 b 乘差分方向的不确定性。取它的上界，就恰好得到算法选择第二个候选时的 UCB。到这里，算法中看似人为加入的探索项，已经有了概率依据。\n\n备用展开：噪声范数不超过 sqrt[2 log(1/δ) + log(det(V_{t−1})/det(V0))]；log 行列式比不超过 d log(1+4κμT/(λd))。这里沿用原稿保守的 1-subgaussian 常数。",
       "原稿 Lemma 1 Step 5、Corollary 1，pp.6–7", "置信界闭合，并解释算法中探索奖励的来源")


# 13 — wrap, 10 sec
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb("navy")
text(s, "把四部分接成一条线", .73, .87, 11.8, .75, 35, "white", True)
text(s, "从一次偏好反馈，走到一次有置信保证的比较。", .77, 1.81, 11.6, .45, 21, "pale")
for i, (label, desc) in enumerate([
    ("问题", "成对比较\n只观察偏好"),
    ("假设", "控制尺度\n保留局部曲率"),
    ("算法", "Ridge-MLE 估计\n贪心 + UCB 选臂"),
    ("置信界", "参数误差界\n推出效用差上界"),
]):
    x = .76 + 3.14*i
    text(s, f"0{i+1}", x, 2.90, 2.65, .39, 16, "66C2C1", True)
    text(s, label, x, 3.53, 2.65, .61, 30, "white", True)
    text(s, desc, x, 4.40, 2.73, .95, 21, "pale")
    if i < 3:
        line(s, x+2.68, 3.84, x+2.99, 3.84, "66C2C1", 2.1, True)
rect(s, .76, 5.78, 11.84, .72, "244057", radius=True)
text(s, "核心连接：置信界给出探索项，探索项指导下一轮比较。", 1.00, 5.98, 11.34, .38,
     22, "white", True)
footer(s, len(prs.slides), 0, "Source: Zhuohua LI · A Clean Proof for Dueling Bandits · §§1–4，pp.1–7", 10, True)
record(s, "把四部分接成一条线", "收束", 10,
       "关键连接是：参数置信界推出效用差置信界，从而给出 UCB 探索项，指导下一轮比较。",
       "原稿 §§1–4，pp.1–7", "模型、假设、算法与置信界形成完整逻辑链")


def save_notes():
    total = sum(item["seconds"] for item in DECK)
    rows = ["# 逐页讲稿与时间安排（v1）", "",
            "内容范围：A Clean Proof for Dueling Bandits，第 1–4 节。",
            "", f"共 {len(DECK)} 页，建议总时长 {total//60} 分 {total%60:02d} 秒。"
            "讲稿也已写入 PPT 的演讲者备注；备注含备用解释，不必逐字念完。", "",
            "| 页码 | 内容 | 建议用时 | 累计 |", "| --- | --- | --- | --- |"]
    elapsed = 0
    for d in DECK:
        elapsed += d["seconds"]
        rows.append(f'| {d["page"]:02d} | {d["title"]} | {d["seconds"]} 秒 | {elapsed//60}:{elapsed%60:02d} |')
    rows += ["", "篇幅分配：问题设定 2:10；假设 2:10；算法 2:30；置信界 2:30；开场、目录与收束共 0:40。",
             "", "讲解深度：MLE 的似然推导、求导和置信界均保留关键步骤；长推导与定理的中间形式写在备注。",
             "", "曲率范围：第 6 页保留原稿算法，并明确置信界需要中间点处的曲率下界；没有把无约束 MLE 悄悄改为球约束估计。", ""]
    for d in DECK:
        rows += [f'## {d["page"]:02d}｜{d["title"]}', "",
                 f'建议用时：{d["seconds"]} 秒。核心信息：{d["key"]}。', "",
                 d["notes"], "", f'来源：{d["source"]}', ""]
    (ROOT / "讲稿与时间安排.md").write_text("\n".join(rows), encoding="utf-8")
    (ROOT / "slide_content.json").write_text(json.dumps(DECK, ensure_ascii=False, indent=2), encoding="utf-8")
    (ASSETS / "equations.json").write_text(json.dumps(FORMULAS, ensure_ascii=False, indent=2), encoding="utf-8")


def audit_geometry():
    problems = []
    for i, slide in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            if shp.left < -Inches(.01) or shp.top < -Inches(.01):
                problems.append({"slide": i, "shape": shp.name, "issue": "negative position"})
            if shp.left+shp.width > prs.slide_width+Inches(.01) or shp.top+shp.height > prs.slide_height+Inches(.01):
                problems.append({"slide": i, "shape": shp.name, "issue": "outside slide"})
    audit = {"slides": len(prs.slides), "duration_seconds": sum(d["seconds"] for d in DECK),
             "notes_present": all(bool(sl.notes_slide.notes_text_frame.text.strip()) for sl in prs.slides),
             "smallest_equation_font_pt": min(f["effective_font_pt"] for f in FORMULAS),
             "geometry_issues": problems}
    (PREVIEW / "structure_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    if problems:
        raise RuntimeError(problems)
    return audit


if __name__ == "__main__":
    # Explicit empty effects prevent inherited Office-theme shadows.
    for slide in prs.slides:
        for shape in slide.shapes:
            sppr = shape._element.find(qn("p:spPr"))
            if sppr is not None:
                for tag in ("a:effectLst", "a:effectDag"):
                    existing = sppr.find(qn(tag))
                    if existing is not None:
                        sppr.remove(existing)
                sppr.append(OxmlElement("a:effectLst"))
    prs.save(str(OUTFILE))
    save_notes()
    audit = audit_geometry()
    print(json.dumps({"presentation": str(OUTFILE), "audit": audit}, ensure_ascii=False, indent=2))
