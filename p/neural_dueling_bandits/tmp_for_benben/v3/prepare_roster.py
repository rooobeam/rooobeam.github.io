"""合并成员表和截图转录，生成可核对的数据；本脚本不生成布局。"""

import csv
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SOURCE = ROOT / "复仇者联盟统计表.xlsx"


def normalize(name):
    name = unicodedata.normalize("NFKC", name)
    name = re.sub(r"^\[[^\]]+\]", "", name)
    name = re.sub(r"\s+", "", name)
    return name.rstrip("。.").casefold()


def is_red(cell):
    color = cell.font.color
    if color is None or color.type != "rgb":
        return False
    rgb = color.rgb[-6:]
    red, green, blue = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    return red >= 150 and red > 1.4 * green and red > 1.4 * blue


def read_members():
    wb = load_workbook(SOURCE, data_only=True)
    ws = wb.active
    blocks = [("FJG", 2), ("VNR", 7), ("VPH", 12), ("OVO", 17), ("KUY", 22), ("sss", 27)]
    members = []
    for group, col in blocks:
        for row in range(3, ws.max_row + 1):
            name_cell = ws.cell(row, col)
            if name_cell.value is None:
                continue
            name = str(name_cell.value).strip()
            if re.fullmatch(r"\d+人", name):
                continue
            # 跨多列的合并文字为说明，不是成员；保留战力空白的真实成员。
            if any(name_cell.coordinate in rng and rng.max_col > col for rng in ws.merged_cells.ranges):
                continue
            power_cell = ws.cell(row, col + 1)
            raw_power = power_cell.value
            if raw_power is not None and not isinstance(raw_power, (int, float)):
                raise ValueError(f"无法识别战力：{power_cell.coordinate}={raw_power!r}")
            members.append({
                "name": name,
                "original_name": name,
                "group": group,
                "power": raw_power,
                "original_power": raw_power,
                "power_source": f"{SOURCE.name}!{ws.title}!{power_cell.coordinate}" if raw_power is not None else None,
                "role": ws.cell(row, col + 2).value,
                "source_cell": name_cell.coordinate,
                "excluded_red": is_red(power_cell),
                "source_order": len(members),
                "screenshot_rank": None,
                "screenshot_name": None,
                "match_status": "无截图排名",
            })
    wb.close()
    if len({normalize(m["name"]) for m in members}) != len(members):
        raise ValueError("成员姓名标准化后有重复，需人工核对。")
    return members


def add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="264B65")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EEF4F7")
        ws.row_dimensions[row[0].row].height = 32
    for col in range(1, len(headers) + 1):
        lengths = [len(str(ws.cell(row, col).value or "")) for row in range(1, ws.max_row + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(58, max(14, max(lengths) * 1.7 + 3))
    return ws


def main():
    config = json.loads((HERE / "ranking_input.json").read_text(encoding="utf-8"))
    members = read_members()
    source_member_count = len(members)
    additional = config.get("additional_members", [])
    for i, entry in enumerate(additional, 1):
        if any(normalize(m["name"]) == normalize(entry["name"]) for m in members):
            raise ValueError(f"新增成员与已有姓名重复：{entry['name']}")
        members.append({
            "name": entry["name"],
            "original_name": None,
            "group": entry.get("group", "待分组"),
            "power": entry.get("power"),
            "original_power": None,
            "power_source": entry["source"],
            "role": None,
            "source_cell": f"USER_ADD_{i}",
            "excluded_red": False,
            "source_order": len(members),
            "screenshot_rank": None,
            "screenshot_name": None,
            "match_status": "无截图排名",
        })
    lookup = {normalize(m["name"]): m for m in members}
    for name, update in config.get("name_overrides", {}).items():
        member = lookup[normalize(name)]
        new_key = normalize(update["name"])
        if new_key in lookup and lookup[new_key] is not member:
            raise ValueError(f"更名后姓名冲突：{update['name']}")
        member["name"] = update["name"]
        member["name_source"] = update["source"]
        lookup[new_key] = member

    ranks = {}
    total_observations = 0
    for screenshot in config["screenshots"]:
        if not (ROOT / screenshot["file"]).is_file():
            raise FileNotFoundError(screenshot["file"])
        for rank, raw_name, stage in screenshot["rows"]:
            total_observations += 1
            if rank in ranks:
                old = ranks[rank]
                if normalize(old["raw_name"]) != normalize(raw_name) or old["stage"] != stage:
                    raise ValueError(f"截图第{rank}名出现冲突")
                if screenshot["file"] not in old["sources"]:
                    old["sources"].append(screenshot["file"])
                continue
            ranks[rank] = {"rank": rank, "raw_name": raw_name, "stage": stage, "sources": [screenshot["file"]]}
    expected_ranks = set(range(1, 56))
    if set(ranks) != expected_ranks:
        raise ValueError(f"截图名次缺失或越界：{sorted(set(ranks) ^ expected_ranks)}")

    aliases = {normalize(a["screenshot_name"]): a for a in config["candidate_aliases"]}
    for rank in sorted(ranks):
        item = ranks[rank]
        key = normalize(item["raw_name"])
        alias = aliases.get(key)
        member = lookup.get(key)
        status = "标准化匹配" if member else "成员表未找到"
        if member is None and alias:
            member = lookup.get(normalize(alias["member_name"]))
            status = "别名已确认" if alias["confirmed"] else "别名暂匹配，待确认"
        item["matched_member"] = member["name"] if member else None
        item["match_status"] = status
        item["transcription_note"] = config.get("transcription_notes", {}).get(str(rank), "")
        if member:
            if member["screenshot_rank"] is not None:
                raise ValueError(f"同一成员匹配多个排名：{member['name']}")
            member["screenshot_rank"] = rank
            member["screenshot_name"] = item["raw_name"]
            member["match_status"] = status

    for name, update in config["power_overrides"].items():
        member = lookup[normalize(name)]
        if update.get("only_if_no_screenshot_rank") and member["screenshot_rank"] is not None:
            continue
        member["power"] = update["power"]
        member["power_source"] = update["source"]

    active = [m for m in members if not m["excluded_red"]]
    excluded = [m for m in members if m["excluded_red"]]

    def priority(m):
        if m["screenshot_rank"] is not None:
            return (0, m["screenshot_rank"], m["source_order"])
        if m["power"] is not None:
            return (1, -m["power"], m["source_order"])
        return (2, 0, m["source_order"])

    active.sort(key=priority)
    resolved = [m for m in active if m["screenshot_rank"] is not None or m["power"] is not None]
    missing = [m for m in active if m["screenshot_rank"] is None and m["power"] is None]
    ranked_without_power = [m for m in active if m["screenshot_rank"] is not None and m["power"] is None]
    for i, member in enumerate(resolved, 1):
        member["provisional_order"] = i
    for member in missing:
        member["provisional_order"] = None

    unmatched = [ranks[r] for r in sorted(ranks) if ranks[r]["matched_member"] is None]
    summary = {
        "source_member_count": source_member_count,
        "additional_member_count": len(additional),
        "red_excluded_count": len(excluded),
        "active_member_count": len(active),
        "user_target_member_count": 100,
        "count_gap_to_target": 100 - len(active),
        "screenshot_observations": total_observations,
        "unique_screenshot_ranks": len(ranks),
        "matched_ranked_members_including_candidate_aliases": sum(m["screenshot_rank"] is not None for m in active),
        "candidate_alias_count": sum(m["match_status"] == "别名暂匹配，待确认" for m in active),
        "unranked_members_with_power": sum(m["screenshot_rank"] is None and m["power"] is not None for m in active),
        "members_missing_both_rank_and_power": len(missing),
        "ranked_members_without_power": len(ranked_without_power),
        "unmatched_screenshot_names": len(unmatched),
        "active_group_counts": dict(Counter(m["group"] for m in active)),
    }
    if len(active) + len(excluded) != len(members) or len(resolved) + len(missing) != len(active):
        raise AssertionError("名单分组未完整覆盖原表和新增成员。")

    report = {
        "status": "数据核对稿；昵称对应待确认；尚未生成正式布局。",
        "sorting_rule": config["sorting_rule"],
        "summary": summary,
        "active_members": active,
        "excluded_members": excluded,
        "screenshot_ranks": [ranks[r] for r in sorted(ranks)],
        "needs_power_or_rank": missing,
        "ranked_without_power": ranked_without_power,
        "unmatched_screenshot_names": unmatched,
        "candidate_aliases": config["candidate_aliases"],
        "name_overrides": config.get("name_overrides", {}),
        "power_overrides": config["power_overrides"],
        "additional_members": additional,
    }
    (HERE / "roster_review.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "核对说明", ["项目", "结果"], [
        ["排序规则", config["sorting_rule"]],
        ["战力补充", "、".join(f"{lookup[normalize(name)]['name']} = {lookup[normalize(name)]['power']}" for name in config["power_overrides"])],
        ["昵称更正", "步青雲已按用户确认更名为平平大王；悲伤小伍保留原字"],
        ["平平大王", "未在已提供的第1～55名截图中找到；按用户指定战力3800排序"],
        ["名单人数", f"原表{source_member_count}人；战力红字剔除{len(excluded)}人；新增{len(additional)}人；最终{len(active)}人"],
        ["新增成员", "、".join(f"{m['name']} = {m.get('power')}" for m in additional) or "无"],
        ["近似姓名", "新增奔波儿灞3000与原表奔霸儿波2900分别记录"],
        ["截图范围", "1～55名连续覆盖；重叠截图和底部固定显示的第20名已去重"],
        ["匹配结果", f"{summary['matched_ranked_members_including_candidate_aliases']}名成员有排名（含{summary['candidate_alias_count']}组待确认别名）"],
        ["仍需补充", f"{len(missing)}人同时缺排名和战力，见“需补排名或战力”"],
        ["无需为排序补战力", f"{len(ranked_without_power)}人战力为空但已有截图排名，直接按排名排"],
        ["截图额外姓名", f"{len(unmatched)}名未在原成员表匹配；单独列出，未自动加入布局名单"],
        ["同战力处理", "未上榜成员若战力相同，保留原表读取顺序；不虚构高低"],
        ["截图右列", "最高到达关卡，不作为战力数值"],
        ["正式制作", "布局图尚未生成；正式制作前先让用户核对预计模样"],
    ])
    member_headers = ["暂排序号", "成员", "原分组", "截图原排名", "战力", "排位依据", "昵称匹配", "表格姓名单元格", "战力来源", "原表姓名"]

    def member_row(m):
        basis = "截图排名" if m["screenshot_rank"] is not None else ("战力" if m["power"] is not None else "待补，暂不编号")
        return [m.get("provisional_order"), m["name"], m["group"], m["screenshot_rank"], m["power"], basis, m["match_status"], m["source_cell"], m["power_source"], m["original_name"]]

    add_sheet(wb, "成员排位核对稿", member_headers, [member_row(m) for m in active])
    add_sheet(wb, "需补排名或战力", ["成员", "原分组", "原表姓名单元格"], [[m["name"], m["group"], m["source_cell"]] for m in missing])
    add_sheet(wb, "有排名无需补战力", ["成员", "截图原排名", "原分组"], [[m["name"], m["screenshot_rank"], m["group"]] for m in ranked_without_power])
    add_sheet(wb, "截图排名1至55", ["原排名", "截图昵称", "关卡（非战力）", "表内对应成员", "匹配状态", "来源图片", "识读说明"], [
        [r["rank"], r["raw_name"], r["stage"], r["matched_member"], r["match_status"], "\n".join(r["sources"]), r["transcription_note"]] for r in report["screenshot_ranks"]
    ])
    add_sheet(wb, "昵称对应待确认", ["截图昵称", "成员表昵称", "状态"], [[a["screenshot_name"], a["member_name"], "已确认" if a["confirmed"] else "暂匹配，待确认"] for a in config["candidate_aliases"]])
    add_sheet(wb, "截图有名但表内未匹配", ["原排名", "截图昵称", "识读说明"], [[r["rank"], r["raw_name"], r["transcription_note"]] for r in unmatched])
    add_sheet(wb, "战力红字淘汰", ["成员", "原分组", "战力", "姓名单元格"], [[m["name"], m["group"], m["power"], m["source_cell"]] for m in excluded])
    add_sheet(wb, "用户补充及更名", ["原表姓名", "采用姓名", "战力", "截图原排名", "补充来源"], [
        [name, lookup[normalize(name)]["name"], lookup[normalize(name)]["power"], lookup[normalize(name)]["screenshot_rank"], update["source"]]
        for name, update in config["power_overrides"].items()
    ])
    add_sheet(wb, "新增成员", ["成员", "战力", "分组", "来源"], [[m["name"], m.get("power"), m.get("group", "待分组"), m["source"]] for m in additional])
    wb.save(HERE / "排名与战力核对.xlsx")

    with (HERE / "成员排位核对稿.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(member_headers)
        writer.writerows(member_row(m) for m in active)

    lines = [
        f"# {HERE.name} 数据核对稿", "",
        "仅整理排名和名单；正式布局须先核对预计模样。", "",
        f"运行：`python {HERE.name}/prepare_roster.py`（需要 openpyxl）。", "",
        config["sorting_rule"], "",
        f"古 塵：战力4100，截图第24名，排位采用截图排名。", "",
        "步青雲已按用户确认更名为平平大王；已提供的第1～55名截图没有此昵称，采用用户指定战力3800。悲伤小伍保留原字，战力3400。", "",
        f"原表{source_member_count}人，红字淘汰{len(excluded)}人，新增{len(additional)}人，最终{len(active)}人。", "",
        f"截图去重后为1～55名，匹配{summary['matched_ranked_members_including_candidate_aliases']}名成员（包含{summary['candidate_alias_count']}组待确认昵称），另有{summary['unranked_members_with_power']}人按已知战力排序。", "",
        "## 用户补充战力", "",
    ]
    lines.extend(f"- {lookup[normalize(name)]['name']}：{lookup[normalize(name)]['power']}" for name in config["power_overrides"])
    lines.extend(["", "## 新增成员", ""])
    lines.extend(f"- {m['name']}：{m.get('power')}；{m['source']}" for m in additional)
    lines.extend(["", "## 仍需排名或战力", ""])
    lines.extend([f"- {m['name']}" for m in missing] or ["无。所有现有成员均有截图排名或已知战力可用于排序。"])
    lines.extend(["", "## 已有排名，无需为排序补战力", ""])
    lines.extend(f"- 第{m['screenshot_rank']}名：{m['name']}" for m in ranked_without_power)
    lines.extend(["", "## 昵称对应待确认", ""])
    lines.extend(f"- {a['screenshot_name']} → {a['member_name']}" for a in config["candidate_aliases"] if not a["confirmed"])
    lines.extend(["", "## 截图名单说明", "", "15名截图上榜者未在成员表匹配，未自动加入名单。第19名昵称有生僻字，暂读“小謾崽”，已单独标注。排名截图中的关卡不等同于战力。", ""])
    (HERE / "核对说明.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print("NEEDS RANK OR POWER:", "、".join(m["name"] for m in missing))
    print("RANKED WITHOUT POWER:", "、".join(f"{m['name']}({m['screenshot_rank']})" for m in ranked_without_power))
    print("TOP 15 MEMBERS:", "、".join(f"{m['name']}[截图{m['screenshot_rank']}]" for m in active[:15]))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
