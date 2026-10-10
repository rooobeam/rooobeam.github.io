"""按已确认的v4格位输出100人正式图和逐人四格坐标，不重新求解格位。"""

import csv
import hashlib
import json
import math
import os
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
os.environ.setdefault("MPLCONFIGDIR", str(HERE / ".mplconfig"))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Polygon, Rectangle
from matplotlib.path import Path as MplPath
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.rcParams["pdf.fonttype"] = 42
FONT = FontProperties(fname="C:/Windows/Fonts/msyh.ttc")
BOLD = FontProperties(fname="C:/Windows/Fonts/msyhbd.ttc")
COLORS = {1: "#b9524f", 2: "#588b75", 3: "#beab6e", 4: "#8a91b1", 5: "#afa3bb"}
TEXT_COLORS = {1: "#ffffff", 2: "#ffffff", 3: "#26352f", 4: "#172738", 5: "#2c283a"}
BACKGROUND = "#f7f5f1"
BOUND = (463, 485, 637, 662)


def check(condition, message):
    if not condition:
        raise ValueError(message)


def name_lines(name):
    if len(name) <= 5 or (name.isascii() and len(name) <= 12):
        return name
    # 以视觉宽度拆行，保留英文姓名完整。
    width = lambda s: sum(1 if ord(c) < 128 else 2 for c in s)
    cuts = [i for i in range(1, len(name)) if not (name[i-1].isascii() and name[i].isascii() and name[i-1].isalnum() and name[i].isalnum())]
    split = min(cuts, key=lambda i: abs(width(name[:i]) - width(name[i:]))) if cuts else len(name) // 2
    return name[:split].strip() + "\n" + name[split:].strip()


def put(ax, x, y, value, size=12, bold=False, **kwargs):
    return ax.text(x, y, value, fontproperties=BOLD if bold else FONT, fontsize=size, **kwargs)


def load_and_validate():
    raw = (HERE / "ring_preview.json").read_bytes()
    approved = json.loads(raw)
    snapshot = json.loads((HERE / "members_snapshot.json").read_text(encoding="utf-8"))
    members = snapshot["active_members"]
    lookup = {m["name"]: m for m in members}
    cx, cy = approved["bear_center"]
    bear = {(x, y) for x in range(cx-1, cx+2) for y in range(cy-1, cy+2)}
    flag_positions = {tuple(p) for p in approved["flag_positions"]}
    flags = []
    for i, (x, y) in enumerate(sorted(flag_positions, key=lambda p: (-p[1], p[0])), 1):
        flags.append({"id": f"F{i:02d}", "x": x, "y": y, "xmin": x-3, "xmax": x+3, "ymin": y-3, "ymax": y+3})
    placements = []
    occupied = {}
    for original in approved["placements_for_review"]:
        row = dict(original)
        x, y = row["anchor"]
        cells = [(x,y), (x+1,y), (x,y+1), (x+1,y+1)]
        row["four_cells"] = cells
        row["group"] = lookup[row["name"]]["group"]
        row["ranking_basis"] = "截图排名" if row["screenshot_rank"] is not None else "战力"
        row["covering_flags"] = [f["id"] for f in flags if all(f["xmin"] <= px <= f["xmax"] and f["ymin"] <= py <= f["ymax"] for px,py in cells)]
        check(row["covering_flags"], f"覆盖不足：{row['name']}")
        k = row["ring"]
        for point in cells:
            px, py = point
            check(463 <= px <= 485 and 637 <= py <= 662, f"越界：{point}")
            check(point not in bear | flag_positions and point not in occupied, f"建筑重叠：{point}")
            check(2*k <= max(abs(px-cx), abs(py-cy)) <= 2*k+1, f"跨圈：{row['name']}")
            occupied[point] = row["name"]
        check(math.isclose(row["distance"], math.hypot(x+.5-cx, y+.5-cy)), "距离记录错误")
        placements.append(row)
    check(len(placements) == 100 and len(occupied) == 400, "人数或城堡占格数错误")
    check([r["name"] for r in placements] == [m["name"] for m in members], "正式版姓名顺序偏离核对名单")
    check([r["order"] for r in placements] == list(range(1,101)), "内部排位不连续")
    check(len({r["name"] for r in placements}) == 100, "姓名重复")
    check([r["distance"] for r in placements] == sorted(r["distance"] for r in placements), "排位与距离不一致")
    anchors = {tuple(r["relative_anchor"]) for r in placements}
    check({(-y-1,x) for x,y in anchors} == anchors, "城堡不满足90度旋转对称")
    relative_flags = {(x-cx,y-cy) for x,y in flag_positions}
    check({(-x,y) for x,y in relative_flags} == relative_flags, "旗子左右不对称")
    check({(x,-y) for x,y in relative_flags} == relative_flags, "旗子上下不对称")
    check({(-y,x) for x,y in relative_flags} == relative_flags, "旗子旋转不对称")
    for flag in flags:
        check(463 <= flag["xmin"] <= flag["xmax"] <= 485 and 637 <= flag["ymin"] <= flag["ymax"] <= 662, "旗子覆盖超出边界")
    counts = Counter(r["ring"] for r in placements)
    check(dict(counts) == {int(k): v for k,v in approved["ring_counts"].items()}, "圈数偏离已确认预览")
    flag_counts = Counter(max(abs(f["x"]-cx),abs(f["y"]-cy))//2 for f in flags)
    stats = []
    for k in range(1,6):
        area = (4*k+3)**2 - (4*k-1)**2
        stats.append({"ring": k, "outer_side": 4*k+3, "area": area, "castles": counts[k], "flags": flag_counts[k], "empty": area - counts[k]*4 - flag_counts[k]})
    checks = [
        ["名单", "通过", "100人、100个姓名无重复，内部排位1～100"],
        ["沿用已确认格位", "通过", "直接读取v4预览，不重新优化或调换格位"],
        ["四格坐标", "通过", "每人4个整数格坐标，共400格"],
        ["无重叠", "通过", "城堡之间、城堡与旗子、城堡与熊坑均不重叠"],
        ["边界", "通过", "所有建筑均在x=463～485、y=637～662内"],
        ["覆盖", "通过", "每人4格均完整处于至少一面旗子的7×7范围内"],
        ["圈宽", "通过", "每圈严格宽2格，每座城堡只属于一个圈"],
        ["城堡对称", "通过", "格位满足90度旋转对称"],
        ["旗子对称", "通过", "同时满足上下镜像、左右镜像、90度旋转对称"],
        ["排位与距离", "通过", "按成员优先级，中心到熊坑中心的距离不递减"],
    ]
    report = {
        "version": "v4", "status": "100人正式版", "approved_by_user": True,
        "approval": "用户：可以，给我100人的版本，也在V4里；还要给个表格，有每个人的四个格的坐标",
        "approved_geometry_sha256": hashlib.sha256(raw).hexdigest(),
        "bounds_inclusive": BOUND, "bear_center": [cx,cy], "bear_cells": sorted(bear),
        "ring_width": 2, "ring_statistics": stats, "flag_positions": flags,
        "cell_order": ["左下格", "右下格", "左上格", "右上格"],
        "coordinate_meaning": "坐标标识格子中心；每座城堡占用所列四个相邻格，城堡几何中心为左下格坐标各加0.5",
        "placements": placements, "validation": checks,
    }
    return report


def header(fig, heading, subtitle):
    fig.text(.045,.957,heading,fontproperties=BOLD,fontsize=29,color="#273943")
    fig.text(.046,.924,subtitle,fontproperties=FONT,fontsize=13,color="#566972")
    fig.text(.954,.959,"V4 · 100人",fontproperties=BOLD,fontsize=15,color="#5b7164",ha="right")


def footer(fig, stats, coordinates=False):
    for i, stat in enumerate(stats):
        x=.047+i*.175
        fig.patches.append(Rectangle((x,.095),.012,.013,facecolor=COLORS[stat["ring"]],transform=fig.transFigure))
        fig.text(x+.018,.095,f"第{stat['ring']}圈  {stat['castles']}人",fontproperties=BOLD,fontsize=12.5,color="#43594f")
    line="城堡小字为左下格坐标；每人的左下、右下、左上、右上四格坐标见配套表格。" if coordinates else "每个姓名块都是2×2城堡；空白为真实空隙，蓝色为1×1旗子，熊坑为3×3。"
    fig.text(.047,.061,line,fontproperties=FONT,fontsize=12,color="#60706c")
    fig.text(.047,.030,"熊坑中心 (474,650)  ·  每圈宽2格  ·  16面旗子  ·  排名优先、其余按战力，排位靠前者分配更近格位",fontproperties=FONT,fontsize=11.5,color="#60706c")


def fit_labels(fig, ax, records):
    # 使用真实字体边界核对100个姓名均在各自城堡内，避免文字跨格。
    fig.canvas.draw()
    for artist, polygon, name in records:
        path = MplPath(ax.transData.transform(polygon))
        for _ in range(8):
            bbox = artist.get_window_extent(fig.canvas.get_renderer()).expanded(1.025,1.03)
            points = [(bbox.x0,bbox.y0),(bbox.x1,bbox.y0),(bbox.x1,bbox.y1),(bbox.x0,bbox.y1)]
            if path.contains_points(points).all():
                break
            artist.set_fontsize(artist.get_fontsize()*.93)
            fig.canvas.draw()
        else:
            raise ValueError(f"姓名超出城堡范围：{name}")
    return len(records)


def top_figure(report):
    fig=plt.figure(figsize=(20,22),facecolor=BACKGROUND)
    header(fig,"熊坑围城 · 100人俯视坐标图","按已确认的V4格位制作｜完整姓名与左下格坐标｜旗子编号对应坐标表")
    ax=fig.add_axes([.075,.15,.88,.74]);ax.set_aspect("equal")
    for x in range(463,486):
        for y in range(637,663):
            ax.add_patch(Rectangle((x-.5,y-.5),1,1,facecolor="#fcfbf7",edgecolor="#d6dcd8",linewidth=.3))
    labels=[]
    for row in report["placements"]:
        x,y=row["anchor"];k=row["ring"]
        ax.add_patch(Rectangle((x-.5,y-.5),2,2,facecolor=COLORS[k],edgecolor="#596960",linewidth=.75,zorder=2))
        for xs,ys in [([x+.5,x+.5],[y-.5,y+1.5]),([x-.5,x+1.5],[y+.5,y+.5])]:
            ax.plot(xs,ys,color="white",alpha=.12,linewidth=.45,zorder=3)
        artist=put(ax,x+.5,y+.72,name_lines(row["name"]),size=13.6,bold=True,ha="center",va="center",color=TEXT_COLORS[k],zorder=7,linespacing=1.08)
        put(ax,x+.5,y-.09,f"({x},{y})",size=9.7,ha="center",va="center",color=TEXT_COLORS[k],alpha=.9,zorder=7)
        labels.append((artist,[(x-.43,y+.22),(x+1.43,y+.22),(x+1.43,y+1.38),(x-.43,y+1.38)],row["name"]))
    cx,cy=report["bear_center"]
    ax.add_patch(Rectangle((cx-1.5,cy-1.5),3,3,facecolor="#e5c87c",edgecolor="#b29855",linewidth=1.1,zorder=4))
    put(ax,cx,cy+.18,"熊坑\n3×3",size=19,bold=True,ha="center",va="center",color="#63502a",zorder=8)
    put(ax,cx,cy-.8,f"({cx},{cy})",size=10,ha="center",va="center",color="#63502a",zorder=8)
    for f in report["flag_positions"]:
        ax.add_patch(Rectangle((f["x"]-.5,f["y"]-.5),1,1,facecolor="#2e638b",edgecolor="#19465f",linewidth=.9,zorder=5))
        put(ax,f["x"],f["y"],f["id"],size=9.5,bold=True,ha="center",va="center",color="white",zorder=8)
    for half in (1.5,3.5,5.5,7.5,9.5,11.5):
        ax.add_patch(Rectangle((cx-half,cy-half),half*2,half*2,fill=False,edgecolor="#6e6d5f",linewidth=.9,linestyle=(0,(4,3)),zorder=6))
    ax.set_xlim(462.45,485.55);ax.set_ylim(636.45,662.55)
    ax.set_xticks(list(range(463,486)));ax.set_yticks(list(range(637,663)))
    ax.tick_params(labelsize=10.5,length=3,color="#75877d")
    ax.set_xlabel("地图 X 坐标",fontproperties=FONT,fontsize=13,labelpad=12,color="#52665b")
    ax.set_ylabel("地图 Y 坐标",fontproperties=FONT,fontsize=13,labelpad=12,color="#52665b")
    for spine in ax.spines.values():spine.set_color("#8c9b93")
    footer(fig,report["ring_statistics"],True)
    count=fit_labels(fig,ax,labels)
    return fig,count


def iso(x,y,z=0):
    return x-y,.57*(x+y)+z


def iso_block(ax,x,y,w,h,color,edge,depth=.06,zorder=2):
    for a,b in [((x,y),(x+w,y)),((x,y),(x,y+h))]:
        points=[iso(*a),iso(*b),iso(*b,depth),iso(*a,depth)]
        ax.add_patch(Polygon(points,facecolor=edge,edgecolor=edge,linewidth=.3,zorder=zorder))
    points=[iso(x,y,depth),iso(x+w,y,depth),iso(x+w,y+h,depth),iso(x,y+h,depth)]
    ax.add_patch(Polygon(points,facecolor=color,edgecolor=edge,linewidth=.55,zorder=zorder+.1))
    return points


def iso_figure(report):
    fig=plt.figure(figsize=(26,18),facecolor=BACKGROUND)
    header(fig,"熊坑围城 · 100人斜视全图","一层城堡一圈，每圈宽2格｜旗子可打断圈｜完整格位呈90°旋转对称")
    ax=fig.add_axes([.025,.15,.95,.72]);ax.set_aspect("equal")
    ax.set_xlim(-27,27);ax.set_ylim(-15.4,15.2);ax.axis("off")
    for x in range(-11,12):
        for y in range(-13,13):
            iso_block(ax,x-.5,y-.5,1,1,"#f1efe9","#d3d7d1",depth=0,zorder=1)
    labels=[]
    for row in report["placements"]:
        x,y=row["relative_anchor"];k=row["ring"]
        polygon=iso_block(ax,x-.5,y-.5,2,2,COLORS[k],"#647369",depth=.06,zorder=3)
        for p1,p2 in [((x+.5,y-.5),(x+.5,y+1.5)),((x-.5,y+.5),(x+1.5,y+.5))]:
            a,b=iso(*p1,.07),iso(*p2,.07)
            ax.plot([a[0],b[0]],[a[1],b[1]],color="white",alpha=.12,linewidth=.35,zorder=4)
        artist=put(ax,*iso(x+.5,y+.5,.09),name_lines(row["name"]),size=12.2,bold=True,ha="center",va="center",color=TEXT_COLORS[k],zorder=20,linespacing=1.05)
        labels.append((artist,polygon,row["name"]))
    iso_block(ax,-1.5,-1.5,3,3,"#e5c87c","#b69b54",depth=.09,zorder=5)
    put(ax,*iso(0,0,.12),"熊坑\n3×3",size=18,bold=True,ha="center",va="center",color="#63502a",zorder=21)
    cx,cy=report["bear_center"]
    for f in sorted(report["flag_positions"],key=lambda a:-(a["x"]+a["y"])):
        x,y=f["x"]-cx,f["y"]-cy
        iso_block(ax,x-.5,y-.5,1,1,"#487ea5","#2c5471",depth=.42,zorder=9)
        put(ax,*iso(x,y,.45),f["id"],size=7.3,bold=True,ha="center",va="center",color="white",zorder=22)
    for x,y,dy in [(-11,-13,-.8),(11,-13,-.85),(-11,12,.85),(11,12,.9)]:
        u,v=iso(x,y)
        put(ax,u,v+dy,f"({cx+x},{cy+y})",size=10.5,ha="center",va="center",color="#6c7a73")
    footer(fig,report["ring_statistics"])
    count=fit_labels(fig,ax,labels)
    return fig,count


def sheet(wb,name,headers,rows,widths=None):
    ws=wb.create_sheet(name);ws.append(headers)
    for row in rows:ws.append(row)
    ws.freeze_panes="C2";ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
    for c in ws[1]:
        c.fill=PatternFill("solid",fgColor="2D5263");c.font=Font(name="微软雅黑",size=11,bold=True,color="FFFFFF")
        c.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=28
        for c in row:
            c.font=Font(name="微软雅黑",size=11,color="253D37")
            c.alignment=Alignment(vertical="center",wrap_text=True)
            if c.row%2==0:c.fill=PatternFill("solid",fgColor="EEF4F1")
    for i,h in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if widths else 18
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows="1:1"
    return ws


def export_tables(report):
    def pair(p):return f"({p[0]}, {p[1]})"
    headers=["内部排位","成员","圈数","左下格(x,y)","右下格(x,y)","左上格(x,y)","右上格(x,y)","截图原排名","战力","原分组","完整覆盖旗子","距熊坑中心"]
    records=[[r["order"],r["name"],r["ring"],*[pair(p) for p in r["four_cells"]],r["screenshot_rank"],r["power"],r["group"]," / ".join(r["covering_flags"]),round(r["distance"],4)] for r in report["placements"]]
    with (HERE/"100人_四格坐标.csv").open("w",encoding="utf-8-sig",newline="") as output:
        writer=csv.writer(output);writer.writerow(headers);writer.writerows(records)
    wb=Workbook();wb.remove(wb.active)
    sheet(wb,"每人四格坐标",headers,records,[11,21,8,19,19,19,19,13,12,13,28,17])
    numeric_headers=["内部排位","成员","左下X","左下Y","右下X","右下Y","左上X","左上Y","右上X","右上Y"]
    sheet(wb,"四格数值便于复制",numeric_headers,[[r["order"],r["name"],*[value for point in r["four_cells"] for value in point]] for r in report["placements"]],[11,21]+[12]*8)
    flag_headers=["编号","旗子X","旗子Y","覆盖最小X","覆盖最大X","覆盖最小Y","覆盖最大Y"]
    flag_rows=[[f["id"],f["x"],f["y"],f["xmin"],f["xmax"],f["ymin"],f["ymax"]] for f in report["flag_positions"]]
    sheet(wb,"旗子坐标与覆盖",flag_headers,flag_rows)
    sheet(wb,"分圈统计",["圈","外框边长","本圈格数","城堡数","旗子数","空白格数"],[[s["ring"],s["outer_side"],s["area"],s["castles"],s["flags"],s["empty"]] for s in report["ring_statistics"]])
    sheet(wb,"校验结果",["检查项","结果","说明"],report["validation"],[26,13,85])
    notes=[
        ["坐标含义",report["coordinate_meaning"]],
        ["四格顺序","左下(x,y)、右下(x+1,y)、左上(x,y+1)、右上(x+1,y+1)，x向右增加，y向上增加"],
        ["示例：刀锋战士","左下(473,652)、右下(474,652)、左上(473,653)、右上(474,653)"],
        ["熊坑","中心(474,650)，占x=473～475、y=649～651共9格"],
        ["可用范围","x=463～485、y=637～662，含两端"],
        ["人数","原表104人，战力红字剔除5人，新增奔波儿灞3000，合计100人"],
        ["同音近似姓名","奔波儿灞3000与奔霸儿波2900是分别记录的两位成员；悲伤小伍保留原字；步青雲采用平平大王"],
        ["排位依据","截图名次优先；其余按战力降序。内部排位不是截图原名次"],
        ["距离与圈数","按实际中心距离安排。部分第五圈正侧位置可能比第四圈转角更近，因此内部排位不总是按圈数连续"],
        ["覆盖旗子","所列每面旗子都能单独完整覆盖该城堡的四格；不表示需要同时依靠所有列出的旗子"],
        ["对称方式","旗子满足上下镜像、左右镜像及90度旋转对称；城堡格位采用90度旋转对称"],
        ["数据来源","沿用已确认的v4/ring_preview.json和members_snapshot.json，没有重新求解或调整已确认格位"],
    ]
    sheet(wb,"摆放说明",["项目","说明"],notes,[25,105])
    ws=wb.create_sheet("地图格子");ws.sheet_view.showGridLines=False;ws.freeze_panes="B2"
    thin=Side(style="thin",color="D4DDD6")
    ws.cell(1,1,"Y / X")
    for x in range(463,486):ws.cell(1,x-463+2,x)
    for y in range(637,663):ws.cell(662-y+2,1,y)
    for row in ws.iter_rows(min_row=1,max_row=27,min_col=1,max_col=24):
        ws.row_dimensions[row[0].row].height=25
        for cell in row:
            cell.font=Font(name="微软雅黑",size=10,color="30453C")
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
            cell.fill=PatternFill("solid",fgColor="F8F7F2")
    ws.column_dimensions["A"].width=9
    for col in range(2,25):ws.column_dimensions[get_column_letter(col)].width=10.5
    def block(x,y,w,h,label,color,foreground):
        left=x-463+2;top=662-(y+h-1)+2
        for row in ws.iter_rows(min_row=top,max_row=top+h-1,min_col=left,max_col=left+w-1):
            for cell in row:cell.fill=PatternFill("solid",fgColor=color.lstrip("#"))
        if w > 1 or h > 1:
            ws.merge_cells(start_row=top,end_row=top+h-1,start_column=left,end_column=left+w-1)
        c=ws.cell(top,left,label);c.font=Font(name="微软雅黑",size=10,bold=True,color=foreground.lstrip("#"))
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for r in report["placements"]:
        x,y=r["anchor"];block(x,y,2,2,r["name"],COLORS[r["ring"]],TEXT_COLORS[r["ring"]])
    cx,cy=report["bear_center"];block(cx-1,cy-1,3,3,"熊坑\n3×3","#E5C87C","#63502A")
    for f in report["flag_positions"]:block(f["x"],f["y"],1,1,f["id"],"#2E638B","#FFFFFF")
    for cell in list(ws[1])+[ws.cell(r,1) for r in range(2,28)]:
        cell.fill=PatternFill("solid",fgColor="2D5263");cell.font=Font(name="微软雅黑",size=10,bold=True,color="FFFFFF")
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="portrait";ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.print_area="A1:X27"
    wb.save(HERE/"100人_四格坐标与旗子.xlsx")


def main():
    report=load_and_validate()
    export_tables(report)
    image_counts={}
    with PdfPages(HERE/"100人_布局打印版.pdf") as pdf:
        pdf.infodict()["Title"]="V4 100人熊坑围城布局"
        for filename,maker in [("100人_斜视全图.png",iso_figure),("100人_俯视坐标图.png",top_figure)]:
            fig,count=maker(report)
            check(count==100,f"{filename}姓名数不为100")
            fig.savefig(HERE/filename,dpi=230,facecolor=fig.get_facecolor())
            pdf.savefig(fig,facecolor=fig.get_facecolor())
            plt.close(fig);image_counts[filename]=count
    report["image_name_counts"]=image_counts
    (HERE/"layout_full.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    (HERE/"validation_full.json").write_text(json.dumps({"status":"passed","castle_count":100,"cell_count":400,"flag_count":16,"image_name_counts":image_counts,"checks":report["validation"]},ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"version":"v4","members":100,"coordinate_cells":400,"flags":16,"images":image_counts,"ring_counts":[s["castles"] for s in report["ring_statistics"]]},ensure_ascii=False))


if __name__=="__main__":
    if hasattr(sys.stdout,"reconfigure"):sys.stdout.reconfigure(encoding="utf-8")
    main()
