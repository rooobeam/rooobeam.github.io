from pathlib import Path
import sys,re,shutil,json,zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=Path(sys.argv[1]).resolve();dest=p.parent
latest=max((p/'exports').glob('*.pptx'),key=lambda f:f.stat().st_mtime)
ppt=dest/'TPOP_v3_EN.pptx';shutil.copy2(latest,ppt)
titles=[
'Beyond generic alignment: whose preference?',
'Problem formulation: personalizing from zero',
'What is personalized, and how is it tested?',
'Two loops: token generation and RM updates',
'Algorithm 1: T-POP',
'Token selection: response 1, then response 2',
'Evidence: stronger alignment, uneven gains',
'Takeaways and limitations']
timing=[45,40,55,95,45,65,50,40]
parts=re.split(r'\n---\s*\n',(p/'notes/total.md').read_text(encoding='utf-8').strip())
assert len(parts)==8
doc=Document();section=doc.sections[0]
section.top_margin=Inches(.7);section.bottom_margin=Inches(.7)
section.left_margin=Inches(.8);section.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.1
for n in ['Title','Heading 1','Heading 2']:
 st=doc.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor.from_string('10243A')
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
footer=section.footer.paragraphs[0];footer.alignment=2;footer.add_run('T-POP v3 · ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
intro='第三版｜8 页全英文页面，中英对照讲稿。主讲选择一种语言，预计约 7–8 分钟。第 4 页展示两个循环，第 5 页保留原算法表，第 6 页解释两次选词的先后关系；实验仅第 7 页。'
md=['# T-POP 第三版｜中英对照讲稿','',intro,''];word_count=0
for i,part in enumerate(parts):
 body='\n'.join(part.strip().splitlines()[1:]).strip()
 paras=[x.strip() for x in re.split(r'\n\s*\n',body) if x.strip()]
 assert len(paras)==2,(i,len(paras))
 zh,en=paras;word_count+=len(re.findall(r"\b[\w’-]+\b",en))
 if i==0:doc.add_heading('T-POP v3 | 中英对照讲稿',0);doc.add_paragraph(intro)
 else:doc.add_page_break()
 heading=f'{i+1:02d}  {titles[i]}';doc.add_heading(heading,1)
 doc.add_paragraph(f'建议用时 / Suggested time: {timing[i]} s')
 doc.add_heading('中文',2);doc.add_paragraph(zh)
 doc.add_heading('English',2);doc.add_paragraph(en)
 md+=['## '+heading,'',f'建议用时：{timing[i]} 秒','','**中文**','',zh,'','**English**','',en,'']
sources=[
'论文事实依据：T-POP: Test-Time Personalization with Online Preference Feedback，arXiv:2509.24696v3；本地 T-POP.pdf。',
'论文链接：https://arxiv.org/abs/2509.24696v3',
'内容组织参考：NICE93-Memory-R1.pptx，第 2–20 页。只借鉴论证顺序，不将其 RAG / memory 内容作为 T-POP 的事实。',
'官方实现：https://github.com/QuZikun/T-POP',
'说明：第 7 页数值是论文 Table 1 四个属性 Average 行的均值。',
'第 5 页图像直接提取自 T-POP.pdf 第 4 页，保留完整 Algorithm 1。',
'第 6 页的 g(v) 只是算法第 12 行梯度差的简写。第 4 页的 δ 对应第 14 行已选延续的梯度差。'
]
doc.add_page_break();doc.add_heading('来源与准备备注 / Sources and preparation notes',1)
for s in sources:doc.add_paragraph(s)
md+=['## 来源与准备备注','']+sources
docx=dest/'TPOP_v3_中英对照讲稿.docx';doc.save(docx)
(dest/'TPOP_v3_中英对照讲稿.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
# Inspect actual PPTX parts, including nested native formulas and embedded notes.
with zipfile.ZipFile(ppt) as z:
 slide_names=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)])
 note_names=[n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)]
 assert len(slide_names)==8 and len(note_names)==8
 ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
 math_count=table_count=picture_count=0
 for n in slide_names:
  rt=ET.fromstring(z.read(n))
  picture_count+=len(rt.findall('.//p:pic',ns))
  table_count+=len(rt.findall('.//a:tbl',ns));math_count+=len(rt.findall('.//m:oMath',ns))
  visible=' '.join(t.text or '' for t in rt.findall('.//a:t',ns))
  assert not re.search('[\u4e00-\u9fff]',visible),n
 for n in note_names:
  t=z.read(n).decode('utf-8')
  assert re.search('[\u4e00-\u9fff]',t) and re.search(r'[A-Za-z]{5}',t),n
 assert table_count==3 and math_count==6 and picture_count==1,(table_count,math_count,picture_count)
receipt={'slides':8,'bilingual_notes':8,'native_tables':table_count,'native_formulas':math_count,'original_algorithm_images':picture_count,'english_words':word_count,'suggested_seconds':sum(timing),'pptx':str(ppt),'docx':str(docx),'postflight':'passed'}
(p/'validation/delivery_check.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False))

