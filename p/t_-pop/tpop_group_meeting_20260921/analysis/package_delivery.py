from pathlib import Path
import sys,re,shutil,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

project=Path(sys.argv[1]).resolve()
dest=project.parent
latest=max((project/'exports').glob('*.pptx'),key=lambda p:p.stat().st_mtime)
pptx=dest/'TPOP_Group_Meeting_EN_8slides.pptx'
shutil.copy2(latest,pptx)
total=(project/'notes/total.md').read_text(encoding='utf-8')
sections=re.split(r'\n---\s*\n',total.strip())
titles=[
'T-POP: one RM, two policies','Text → vector → personal reward',
'Train on complete response pairs','RM reward changes the next token',
'Shared candidates, separate prefixes','Explore uncertain preference differences',
'Two clocks: update V, then train RM','Deployment returns one answer']
durations=[45,50,60,50,55,60,55,40]
intro='8 页英文 PPT；逐页中英对照。选一种语言主讲，预计约 7 分钟；另一种用于理解或备用，不需要两种语言都念。'
md=['# T-POP 组会讲稿｜Chinese–English Script','',intro,'']
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
normal=doc.styles['Normal']
normal.font.name='Arial';normal.font.size=Pt(11)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.line_spacing=1.12
for name in ['Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor.from_string('10243A')
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
footer=sec.footer.paragraphs[0]
footer.alignment=2
footer.add_run('T-POP · ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
words=0
for i,s in enumerate(sections):
 lines=s.strip().splitlines()
 body='\n'.join(lines[1:]).strip()
 paragraphs=[x.strip() for x in re.split(r'\n\s*\n',body) if x.strip()]
 assert len(paragraphs)==2,(i,len(paragraphs))
 zh,en=paragraphs
 words+=len(re.findall(r"\b[\w’-]+\b",en))
 if i==0:
  doc.add_heading('T-POP | 中英对照讲稿',0)
  doc.add_paragraph(intro)
 else:doc.add_page_break()
 heading=f'{i+1:02d}  {titles[i]}'
 doc.add_heading(heading,1)
 timing=f'建议用时 / Suggested time: {durations[i]} s'
 doc.add_paragraph(timing).runs[0].font.color.rgb=RGBColor.from_string('5B6B7A')
 doc.add_heading('中文',2);doc.add_paragraph(zh)
 doc.add_heading('English',2);doc.add_paragraph(en)
 md+=['## '+heading,'',timing,'','**中文**','',zh,'','**English**','',en,'']
doc.add_page_break();doc.add_heading('简短问答与来源 / Q&A and sources',1)
qas=[
('为什么有两条臂，却只有一个 RM？ / Why two arms but one RM?',
 '“臂”是选词策略。两条策略共享模型参数，通过不同选词规则产生比较样本。',
 'An arm is a token-selection policy. Shared model parameters support two policies that generate preference comparisons.'),
('V 更新算不算训练？ / Does updating V train the RM?',
 '不改变 theta。它累积已选延续的梯度差外积，为后续探索估计不确定性。以 V=I、delta=(1,0) 为例，更新后 V=diag(2,1)，同方向不确定性从 1 降为 1/sqrt(2)。',
 'It does not change theta. It accumulates gradient-difference outer products for uncertainty. For V=I and delta=(1,0), V becomes diag(2,1), reducing uncertainty in that direction from 1 to 1/sqrt(2).'),
('训练和推理中的批量维度相同吗？ / Are the batch axes the same?',
 '训练中一批 b 对偏好对应 2b 条文本；推理中 C 表示某条分支当前的候选延续数量。第 2 页的 B 只是进入编码器的文本总数。',
 'A training batch of b preference pairs contains 2b texts. At inference, C counts candidate continuations for one arm. B on slide 2 simply counts encoder input texts.'),
('读公式时应注意什么？ / Which formulation does the talk use?',
 '主讲遵循论文：概率加原始 RM 奖励。官方代码对奖励使用 sigmoid 后再合并，复现时需要区分。完整回答偏好用于前缀评分，其迁移可靠性可以作为讨论点。',
 'The talk follows the paper: probability plus raw RM reward. The public code applies sigmoid to the reward before combining scores. Also, transferring complete-answer supervision to prefix scoring is a useful discussion point.')
]
md+=['## 简短问答与来源 / Q&A and sources','']
for question,zh,en in qas:
 doc.add_heading(question,2);doc.add_paragraph(zh);doc.add_paragraph(en)
 md+=['### '+question,'',zh,'',en,'']
sources=[
'T-POP: Test-Time Personalization with Online Preference Feedback — arXiv:2509.24696v2. https://arxiv.org/abs/2509.24696v2',
'Official implementation: https://github.com/QuZikun/T-POP',
'Presentation workflow: ppt-master. https://github.com/hugohe3/ppt-master'
]
doc.add_heading('Sources',2)
for s in sources:doc.add_paragraph(s)
md+=sources
docpath=dest/'TPOP_Speaker_Notes_ZH_EN.docx'
mdpath=dest/'TPOP_Speaker_Notes_ZH_EN.md'
doc.save(docpath);mdpath.write_text('\n'.join(md)+'\n',encoding='utf-8')
print(json.dumps({'pptx':str(pptx),'docx':str(docpath),'markdown':str(mdpath),'slides':8,'english_words':words,'suggested_seconds':sum(durations)},ensure_ascii=False))

